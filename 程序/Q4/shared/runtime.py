"""Shared cumulative small-test budget and frozen inputs. No annual runner."""
import os,sys,time,json,gzip,hashlib,csv,copy,math
from pathlib import Path
from datetime import datetime,timedelta
import numpy as np
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent; ROOT=next(p for p in Path(__file__).resolve().parents if (p/'C题'/'附件'/'附件1.xlsx').is_file())
sys.path.insert(0,str(HERE/'vendor'))
from core import serial,digest
from data_access import minutes,Access
from forecasting import array_hash,base_predict
STATE=HERE/'work'; STATE.mkdir(exist_ok=True)
OLD=HERE.parent/'closure_tests/outputs/20260913_113753_029830'
def encoded(x): return json.dumps(x,ensure_ascii=False,sort_keys=True,default=serial,allow_nan=False,separators=(',',':')).encode('utf-8')
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(p):
    p=Path(p); raw=p.read_bytes(); return json.loads(gzip.decompress(raw) if p.suffix=='.gz' else raw)
def write(p,x):
    p=Path(p); p.parent.mkdir(exist_ok=True,parents=True); raw=encoded(x)
    if p.suffix=='.gz': raw=gzip.compress(raw,compresslevel=1,mtime=0)
    tmp=p.with_name(p.name+'.tmp')
    with tmp.open('wb') as f: f.write(raw); f.flush(); os.fsync(f.fileno())
    os.replace(tmp,p); return sha(p)
class Budget:
    def __init__(self,stage,seconds=110):
        self.stage=stage; self.out=STATE/stage; self.out.mkdir(exist_ok=True); self.start=time.perf_counter(); self.seconds=seconds; self.timings={}
        base=read(OLD/'report.json'); proc=read(OLD/'process_exit.json')
        self.prior_calls=base['budget']['calls']; self.prior_fits=base['budget']['fits']; self.prior_wall=proc['process_wall_seconds']
        self.events=STATE/'events.jsonl'; self.calls=0; self.fits=0
        if self.events.exists():
            for e in self.events.read_text(encoding='utf-8').splitlines():
                x=json.loads(e); self.calls+=x['kind']=='lp'; self.fits+=x['kind']=='fit'
        self.completed_wall=sum(read(p)['wall_seconds'] for p in STATE.glob('stage_*.json'))
    def remaining(self): return min(self.seconds-(time.perf_counter()-self.start),300-self.prior_wall-self.completed_wall-(time.perf_counter()-self.start))
    def check(self):
        if self.remaining()<=0: raise RuntimeError('small-test cumulative/stage deadline reached')
    def reserve(self,kind,label):
        self.check(); kind='fit' if kind=='fit' else 'lp'
        if kind=='fit':
            if self.prior_fits+self.fits>=40: raise RuntimeError('40 cumulative fits exhausted')
            self.fits+=1
        else:
            if self.prior_calls+self.calls>=1000: raise RuntimeError('1000 cumulative solves exhausted')
            self.calls+=1
        with self.events.open('ab') as f:
            f.write(encoded(dict(stage=self.stage,kind=kind,label=label,elapsed=time.perf_counter()-self.start))+b'\n'); f.flush(); os.fsync(f.fileno())
    def addtime(self,k,s): self.timings[k]=self.timings.get(k,0)+s
    def record(self): return dict(stage=self.stage,wall_seconds=time.perf_counter()-self.start,cumulative_calls=self.calls+self.prior_calls,cumulative_fits=self.fits+self.prior_fits,cumulative_wall_seconds=self.prior_wall+self.completed_wall+time.perf_counter()-self.start,timings=self.timings)
class Sources:
    def __init__(self): self.hashes={}
    def stamp(self,p):
        p=Path(p); self.hashes[str(p)]=sha(p); return p
    def verify(self):
        for p,h in self.hashes.items():
            if sha(p)!=h: raise ValueError('source mutated '+p)
def sheet(src,name):
    wb=load_workbook(src,read_only=True,data_only=True); rows=list(wb[name].values); wb.close()
    assert [minutes(v) for v in rows[0][1:]]==list(range(10,1441,10))
    assert [(r[0].date()-datetime(2025,1,1).date()).days for r in rows[1:]]==list(range(365))
    a=np.array([r[1:] for r in rows[1:]],float).ravel(); assert a.shape==(52560,) and np.isfinite(a).all(); return a
def load_inputs(sources):
    a2=sources.stamp(ROOT/'C题/附件/附件2.xlsx'); a4=sources.stamp(ROOT/'C题/附件/附件4.xlsx'); a3=sources.stamp(ROOT/'C题/附件/附件3.xlsx')
    load=sheet(a2,'小区负载'); pv=sheet(a2,'光伏发电实际功率'); price=sheet(a4,'Sheet1')
    assert min(load)>=0 and min(pv)>=0 and min(price)>0
    wb=load_workbook(a3,read_only=True,data_only=True); rows=list(wb.active.values); wb.close(); official={}; day=None
    for row in rows[1:]:
        if row[0] is not None and str(row[0]).strip(): day=(datetime.strptime(str(row[0]).split()[0],'%Y-%m-%d')-datetime(2025,1,1)).days
        key=(day,minutes(row[1])//10); a=np.array(row[2:],float)
        assert key not in official and a.shape==(24,) and np.isfinite(a).all() and min(a)>=0
        official[key]=a
    assert set(official)=={(d,s) for d in range(365) for s in (0,36,72,108)}
    return dict(load_nodes=load,pv_nodes=pv,price_nodes=price,official=official,prices=np.ones(144))
class View(Access):
    def __init__(self,data,mapping,policy='C2',position=0):
        super().__init__(data,policy,position); self.mapping=mapping
        self.__raw={v:np.asarray(data[k]).copy() for v,k in [('load','load_nodes'),('pv','pv_nodes'),('price','price_nodes')]}
    def completed(self,end,variable):
        if not 0<=end<=self.position: raise ValueError('future completed history denied')
        ids=np.arange(end) if self.mapping=='Q42' else np.maximum(np.arange(end)-1,0)
        return self.__raw[variable][ids].copy()
    def observations_completed_by(self,end,variable):
        a=self.completed(end,variable); self._event('completed',end=end,variable=variable); return a
    def history_before(self,day,variable): return self.completed(day*144,variable).reshape(day,144)
    def root_net(self):
        if self.mapping!='Q42' or self.committed is None: raise ValueError('R2 current net before fixed commitment')
        i=self.position; self._event('root_net',source=i); return float((self.__raw['load'][i]-self.__raw['pv'][i])/6)
    def delivered(self):
        if self.committed is None: raise ValueError('actual settlement before commitment')
        j=self.position if self.mapping=='Q42' else max(self.position-1,0)
        self._event('delivery',source=j); return tuple(float(self.__raw[k][j]) for k in ('load','pv','price'))
class Bases:
    def __init__(self,data,sources):
        self.sources=sources; self.data=data; self.hot={}; self.refs={}; self.pool={}; self.delta={}; self.validation=[]
    def get(self,view,day):
        if view.position<day*144: raise ValueError('future base publication')
        key=(view.mapping,day); hist={v:view.history_before(day,v) for v in ('load','pv')}
        hh={v:array_hash(x) for v,x in hist.items()}
        if key in self.hot:
            if self.refs[key]['histories']!=hh: raise ValueError('cache completed history mismatch')
            return [x.copy() for x in self.hot[key]]
        if day==1:
            vals=[np.tile(hist[v][-1],2) for v in ('load','pv')]; meta={'histories':hh,'method':'original short-history previous day','day':day,'mapping':view.mapping}
        elif view.mapping=='Q42':
            path=self.sources.stamp(ROOT/f'结果/Q2/R2/forecasts/{day:03d}.json'); obj=read(path)
            manifest=read(self.sources.stamp(path.parent.parent/'manifest.json'))
            assert next(v for k,v in manifest['inputs'].items() if k.replace('\\','/').endswith('C题/附件/附件2.xlsx'))==self.sources.hashes[str(ROOT/'C题/附件/附件2.xlsx')]
            for ref in obj['refs']:
                k=ref['k']; v=ref['series']; a=np.array(obj['predictions'][str(k)][v],float)
                assert ref['publish_day']==day and ref['target_day']==day+k and ref['training_target_end_inclusive']<day
                assert ref['forecast_sha256']==hashlib.sha256(a.tobytes()).hexdigest()
            vals=[np.concatenate([obj['predictions'][str(k)][v] for k in (0,1)]) for v in ('load','pv')]
            meta={'histories':hh,'path':str(path),'hash':sha(path),'refs':obj['refs'],'day':day,'mapping':view.mapping}
        else:
            base=ROOT/'结果/Q3'
            commit=read(self.sources.stamp(base/f'commits/{day:03d}.json')); key3=commit['states']['C2']['state']['base_id']
            path=self.sources.stamp(base/f'cache/{key3}.json.gz'); obj=read(path)
            assert obj['day']==day and obj['history_cutoff']==day*144 and obj['history_hashes']==hh and obj['key']==key3
            assert obj['identity']['prediction_sha256']==sha(HERE/'vendor/forecasting.py')
            assert obj['identity']['source_sha256']['C题/附件/附件2.xlsx']==self.sources.hashes[str(ROOT/'C题/附件/附件2.xlsx')]
            for task in obj['tasks']:
                assert task['issued_position']==task['history_cutoff']==day*144 and max(task['record']['target_days'])<day
            vals=[np.array(obj['predictions'][v],float) for v in ('load','pv')]
            meta={'histories':hh,'path':str(path),'hash':sha(path),'tasks':obj['tasks'],'day':day,'mapping':view.mapping}
        assert all(x.shape==(288,) and np.isfinite(x).all() and min(x)>=0 for x in vals)
        self.hot[key]=vals; self.refs[key]=meta; return [x.copy() for x in vals]
    def risk(self,view,day):
        if not self.pool:
            root=ROOT/'结果/Q2/M0'
            p=self.sources.stamp(root/'forecast_archive.csv')
            with p.open(encoding='utf-8-sig',newline='') as f:
                for row in csv.DictReader(f):
                    d=int(row['publish_day_index']); t=int(row['slot'])-1
                    if 21<=d<=31:
                        self.pool.setdefault(d,np.zeros(144))[t]=float(row['net_hat_kwh'])
                        assert abs((float(row['load_hat_kw'])-float(row['pv_hat_kw']))/6-float(row['net_hat_kwh']))<1e-7
            with self.sources.stamp(root/'delta_pool.csv').open(encoding='utf-8-sig',newline='') as f:
                for row in csv.reader(f):
                    if row[0] in ('31','32'): self.delta[int(row[0])]=np.array(row[2:],float)
        past=(view.history_before(day,'load')-view.history_before(day,'pv'))/6
        errors=np.array([past[d]-self.pool[d] for d in range(21,day)])
        expected=[]
        for t in range(144):
            pool=sorted(float(x) for x in errors[:,max(0,t-2):min(144,t+3)].ravel()); z=.8*(len(pool)-1); k=math.floor(z); f=z-k
            expected.append(pool[k]*(1-f)+pool[min(k+1,len(pool)-1)]*f)
        err=float(max(abs(np.array(expected)-self.delta[day]))); assert err<1e-6
        self.validation.append({'kind':'M0 saved delta vs independently reconstructed rounded pool','day':day,'max_difference_kwh':err,'cache_precision':'12 significant digits, verified within 1e-6 kWh; not byte-identical unrounded historical coefficients'})
        return self.delta[day].copy()
class Prices:
    def __init__(self,budget):
        self.budget=budget; self.path=STATE/'price_cache.json'; self.coeffs=read(self.path) if self.path.exists() else {}; self.records=[]
        for c in self.coeffs.values(): self.validate_coefficient(c)
    @staticmethod
    def validate_coefficient(c):
        fields=('a','b','rho','rho_raw')
        if not all(k in c and np.isfinite(c[k]) for k in fields) or not 0<=c['rho']<=1: raise ValueError('invalid price coefficient')
        if c.get('content_hash')!=digest({k:v for k,v in c.items() if k!='content_hash'}): raise ValueError('price coefficient integrity mismatch')
    def coefficient(self,view,day):
        r0=day*144; p=view.completed(r0,'price'); n=(view.completed(r0,'load')-view.completed(r0,'pv'))/1000
        if not np.isfinite(p).all() or not np.isfinite(n).all(): raise ValueError('invalid actual history')
        ids=np.arange(max(1008,r0-4032),r0)
        if len(ids)<1008: return None
        key=f'{view.mapping}:{day}'; history=digest([p,n])
        if key in self.coeffs:
            c=self.coeffs[key]
            self.validate_coefficient(c)
            if c['history']!=history or c['day']!=day or c['last']!=r0-1 or c['first']!=max(1008,r0-4032): raise ValueError('price coefficient history identity mismatch')
            return c
        x=n[ids]-n[ids-1008]; y=p[ids]-p[ids-1008]; dx=x-x.mean(); den=dx@dx
        self.budget.reserve('fit',key+'_OLS'); b=float(dx@(y-y.mean())/den) if den>1e-12 else 0.; a=float(y.mean()-b*x.mean()); e=y-a-b*x
        self.budget.reserve('fit',key+'_AR'); denrho=e[:-1]@e[:-1]; raw=float(e[:-1]@e[1:]/denrho) if denrho>1e-12 else 0.
        c=dict(a=a,b=b,rho=float(np.clip(raw,0,1)),rho_raw=raw,history=history,day=day,first=int(ids[0]),last=int(ids[-1]),rows=len(ids),degenerate_b=bool(den<=1e-12),degenerate_rho=bool(denrho<=1e-12))
        c['content_hash']=digest(c)
        self.coeffs[key]=c; write(self.path,self.coeffs); return c
    def predict(self,view,targets,nhat,method,source):
        r=view.position; targets=np.asarray(targets,int); nhat=np.asarray(nhat,float)
        if method not in ('P1','P7') or len(targets)!=len(nhat) or not np.isfinite(nhat).all(): raise ValueError('invalid forecast')
        if np.any(targets<r) or np.any(targets>=min((r//144+2)*144,52560)): raise ValueError('forecast horizon outside law')
        p=view.completed(r,'price'); n=(view.completed(r,'load')-view.completed(r,'pv'))/1000
        if not np.isfinite(p).all() or not np.isfinite(n).all() or np.any(p<=0): raise ValueError('invalid completed price/net history')
        c=self.coefficient(view,r//144) if method=='P1' else None; vals=[]; modes=[]; src=[]; floor=0
        for i,N in zip(targets,nhat):
            lag=int(i-1008)
            if c is not None:
                if not 0<=lag<r: raise ValueError('illegal week source')
                residual=p[-1]-p[-1009]-c['a']-c['b']*(n[-1]-n[-1009])
                v=p[lag]+c['a']+c['b']*(N-n[lag])+c['rho']**int(i-r+1)*residual
                floor+=v<1e-6; vals.append(max(1e-6,float(v))); modes.append('P1'); src.append([lag,r-1,r-1009])
            elif 0<=lag<r: vals.append(float(p[lag])); modes.append('P7'); src.append([lag])
            else:
                idx=list(range(int(i%144),r,144))[-7:] or list(range(max(0,r-144),r))
                if not idx: raise ValueError('no price history: do not optimize Jan1 midnight')
                vals.append(float(np.mean(p[idx]))); modes.append('P0'); src.append(idx)
        record=dict(mapping=view.mapping,r=r,targets=targets,nhat_mw=nhat,source=source,method=method,coefficient=c,modes=modes,actual_sources=src,floor_count=int(floor),prices=np.array(vals),history_last=r-1)
        self.records.append(record); return np.array(vals),record
