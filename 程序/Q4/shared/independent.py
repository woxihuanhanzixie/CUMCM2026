"""Raw ledger/contract/formula audit; imports no controller or optimizer."""
import math,copy
import numpy as np
from openpyxl import load_workbook
from runtime import ROOT,read,write,digest
def raw_sources():
    out={}
    for rel,sheet,key in [('附件2.xlsx','小区负载','load'),('附件2.xlsx','光伏发电实际功率','pv'),('附件4.xlsx','Sheet1','price')]:
        wb=load_workbook(ROOT/'C题/附件'/rel,read_only=True,data_only=True)
        out[key]=np.array([[float(x) for x in row] for row in wb[sheet].iter_rows(min_row=2,max_row=366,min_col=2,max_col=145,values_only=True)]).ravel(); wb.close()
    return out
def price_check(rec,raw):
    r=rec['r']; targets=np.array(rec['targets']); mapping=rec['mapping']
    # Read physical intervals from raw nodes independently of View.
    indices=np.arange(r) if mapping=='Q42' else np.array([max(k-1,0) for k in range(r)])
    p=raw['price'][indices]; n=(raw['load'][indices]-raw['pv'][indices])/1000; c=rec['coefficient']; vals=[]
    for i,N,mode,used in zip(targets,rec['nhat_mw'],rec['modes'],rec['actual_sources']):
        assert all(0<=s<r for s in used)
        if mode=='P1':
            assert c is not None and c['day']==r//144 and c['last']==c['day']*144-1 and c['first']==max(1008,c['day']*144-4032)
            e=p[r-1]-p[r-1009]-c['a']-c['b']*(n[r-1]-n[r-1009])
            value=max(1e-6,p[i-1008]+c['a']+c['b']*(N-n[i-1008])+c['rho']**int(i-r+1)*e)
            assert used==[int(i-1008),r-1,r-1009]
        elif mode=='P7': value=p[i-1008]; assert used==[int(i-1008)]
        else:
            assert mode=='P0'; ids=[j for j in range(r) if j%144==i%144][-7:] or list(range(max(0,r-144),r)); assert used==ids; value=math.fsum(float(p[j]) for j in ids)/len(ids)
        vals.append(value)
    err=float(np.max(abs(np.array(vals)-rec['prices']))); assert err<1e-10; return err
def audit(bundle,raw,expected_segments=288):
    rows=bundle['ledger']; plans=bundle['plans']; versions=bundle['versions']; start=bundle['start']; name=bundle['name']; q3=name.startswith('C'); mode='Q43' if q3 else 'Q42'
    assert len(rows)==expected_segments and [x['position'] for x in rows]==list(range(start,start+expected_segments))
    E=6000.; max_phys=max_bill=max_price=max_obj=0.; bills=[]; failures=0; unproven=0; secondary=0
    expected_clocks=list(range(max(start,36),start+expected_segments,6)) if q3 else ([k for k in range(start,start+expected_segments) if k%144==0] if name=='M0' else None)
    if expected_clocks is not None: assert [p['position'] for p in plans]==expected_clocks
    elif expected_segments==288: assert len(plans)==290 and sum(x['stage']=='root' for x in plans)==288
    if expected_segments==288:
        allowed=[k for k in range(start,start+288) if k%144==0 or (name=='C2' and k%144 in (36,72,108))]
        assert [v['published_position'] for v in versions]==allowed
    by_hour={p['position']:p for p in plans if q3}
    for row in rows:
        i=row['position']; d,t=divmod(i,144); node=i if not q3 else max(i-1,0)
        ell=float(raw['load'][node]/6); pv=float(raw['pv'][node]/6); p=float(raw['price'][node])
        assert max(abs(row['ell']-ell),abs(row['pv']-pv),abs(row['price']-p),abs(row['E_before']-E))<1e-6
        vv=[v for v in versions if v['start']<=i<v['end'] and v['published_position']<=i]; final=vv[-1]; orig=[v for v in vv if v['kind']=='original'][0]
        a=float(final['orders'][i-final['start']]); q=float(orig['orders'][t])
        assert abs(a-row['a'])<1e-6 and abs(q-row['q0'])<1e-6 and row['version_id']==final['id'] and row['published']==final['published_position']
        if name!='C2': assert a==q
        c,s,g,w=(float(row[k]) for k in ('c','s','g','w')); nextE=E+.9*c-s/.9
        err=max(abs(a+pv+s+g-ell-c-w),abs(row['E_after']-nextE),max(0,1200-nextE,nextE-10800,c-5000/6,s-5000/6,-min(a,c,s,g,w)))
        assert err<1e-6 and all(min(x,y)<1e-6 for x,y in ((c,s),(c,g),(s,w),(g,w))); max_phys=max(max_phys,err)
        if q3:
            R=1200. if i<36 else by_hour[i-i%6]['solution']['E'][i%6]; assert abs(R-row['R'])<1e-6
            net=a+pv-ell
            if net>=0: ec=min(net,5000/6,(10800-E)/.9); es=eg=0.; ew=net-ec
            else: ec=ew=0.; es=min(-net,5000/6,.9*max(E-R,0)); eg=-net-es
            assert max(abs(c-ec),abs(s-es),abs(g-eg),abs(w-ew))<1e-6
        elif name=='M0':
            net=a+pv-ell; ec=min(max(net,0),5000/6,(10800-E)/.9); es=min(max(-net,0),5000/6,.9*(E-1200)); assert max(abs(ec-c),abs(es-s))<1e-6
        normal=p*a; adjust=(p*min(q,a)+.5*p*max(q-a,0)+1.5*p*max(a-q,0)-normal) if q3 else 0.; emergency=5*p*g; total=normal+adjust+emergency
        max_bill=max(max_bill,abs(total-row['total']),abs(normal-row['normal']),abs(adjust-row['adjustment']),abs(emergency-row['emergency'])); bills.append(total); E=nextE
    assert max_bill<1e-6 and abs(math.fsum(bills)-math.fsum(r['total'] for r in rows))<.01
    for plan in plans:
        rec=plan['price_record']; max_price=max(max_price,price_check(rec,raw)); i=plan['position']; t=i%144
        expected_length=144 if name=='M0' else 288-t
        assert len(rec['prices'])==expected_length and rec['r']==i
        sol=plan['solution']; p=np.array(rec['prices'])
        if q3:
            problem=plan['problem']; A=np.array(sol['A']); u=np.array(sol['u']); en=np.array(sol['E']); ell=np.array(problem['ell']); pv=np.array(problem['pv']); n=ell-pv
            assert np.max(abs(np.diff(np.r_[problem['energy'],en])-u))<1e-6
            assert max(np.max(1200-en),np.max(en-10800))<1e-6
            if plan['stage']=='feedback': assert np.max(abs(A[:144-t]-problem['current']))<1e-6
            c=np.maximum(u,0)/.9; s=np.maximum(-u,0)*.9; g=np.maximum(n-A+c-s,0)
            adjustment=np.zeros(len(A))
            if plan['stage']!='original': adjustment[:144-t]=.5*abs(A[:144-t]-problem['q0'])
            fee=float(p@(A+adjustment+5*g)); max_obj=max(max_obj,abs(fee-sol['objective'])); assert abs(fee-sol['objective'])<1e-6 and abs(fee-sol['lower'])<1e-5
            assert sol['residual']<1e-6 and sol['stationarity']<1e-7
            assert np.max(abs(np.array(rec['nhat_mw'])-(ell-pv)*6/1000))<1e-12
        elif not sol['usable']: failures+=1
        else:
            unproven+=sol['status']!=0; n=np.array(plan['nbar']); key='q' if name=='M0' else 'Q'; A=np.array(sol[key]); c=np.array(sol['c']); s=np.array(sol['s']); w=np.array(sol['w']); en=np.array(sol['E']); g=np.array(sol.get('z',np.zeros(len(A))))
            assert np.max(abs(A-c+s+g-w-n))<1e-6 and np.max(abs(np.diff(np.r_[plan['energy'],en])-.9*c+s/.9))<1e-6
            if name=='M0': assert np.max(w-np.array(plan['wcap']))<1e-6
            if plan['stage']=='root':
                nf=len(plan['q_fixed']); fee=float(5*p[:nf]@g[:nf]+p[nf:]@A[nf:]); assert np.max(abs(A[:nf]-plan['q_fixed']))<1e-6
                sec=sol['secondary']; secondary+=sec['accepted']
                if sec['accepted']: assert sec['main_cost_yuan']<=sol['primary_fun']+1.1e-6 and sec['root_energy_gain_kwh']>=-1e-6
                rr=plan['restored']; assert rr['obj_restored']<=fee+1e-6
                raw_net=(raw['load'][i]-raw['pv'][i])/6; assert abs(n[0]-raw_net)<1e-6
            else: fee=float(p@A)
            max_obj=max(max_obj,abs(fee-sol['fun'])); assert abs(fee-sol['fun'])<1e-6
    # Authoritative legal order times and price sources; no future actual reads in prediction.
    for e in bundle['access_log']:
        if e['kind']=='completed': assert e['end']<=e['position']
        if e['kind']=='official': assert e['day']*144+e['slot']<=e['position'] and (name!='C0' or e['slot']==0)
    for corr in bundle['corrections']:
        day,slot=divmod(corr['r'],144); shifts=np.asarray(corr['indices'])
        positions=corr['r']+shifts; nodes=positions if not q3 else np.maximum(positions-1,0)
        actual_bound=float(1.5*np.max(raw['price'][nodes])*(-np.asarray(corr['raw'])).sum()); assert actual_bound<=1e-6
    return dict(status='PASS',segments=len(rows),plans=len(plans),normal=math.fsum(r['normal'] for r in rows),adjustment=math.fsum(r['adjustment'] for r in rows),emergency=math.fsum(r['emergency'] for r in rows),total=math.fsum(bills),initial_energy=6000.,final_energy=E,max_physical_residual=max_phys,max_bill_error=max_bill,max_price_rebuild_error=max_price,max_plan_objective_error=max_obj,fallbacks=failures,unproven=unproven,secondary_accepted=secondary,numerical_zero_events=len(bundle['corrections']))
def export_boundaries(bundles):
    checks=[]
    for bundle in bundles:
        start=bundle['start']; rows={r['position']:r for r in bundle['ledger']}
        if start!=4464: continue
        # Template first row is Feb1 00:10 through Feb2 00:00; EO must join the next day.
        window=[rows[i] for i in range(4465,4609)]
        assert window[-1]['position']==4608 and window[-1]['published']==4608
        checks.append(dict(branch=bundle['name']+'-'+bundle['method'],EO_position=4608,EO_order=window[-1]['a'],EO_published=4608,window_total=math.fsum(r['a'] for r in window),window_cost=math.fsum(r['total'] for r in window),Feb1_first_segment_in_natural_ledger=4464 in rows))
    # No 2026 extrapolation: the missing contract and dependent totals remain explicit nulls.
    end={'EO':None,'EP':None,'EQ':None,'reason':'2026 formal contract absent; Q42 also lacks realized endpoint price'}
    assert all(end[k] is None for k in ('EO','EP','EQ'))
    return dict(status='PASS_MAPPING_FIXTURES',cross_day=checks,year_end=end,formal_workbooks_created=False)
