"""Template-preserving export with a separate workbook readback from daily contracts."""
import csv,math,json
from datetime import datetime,timedelta
from pathlib import Path
from openpyxl import load_workbook
import subprocess
from annual import ROOT,HERE,BRANCHES,read,save,unpack,sha

BASE=datetime(2025,1,1)
FIELDS=['position','delivery_start','delivery_end','q0','a_final','c','s','z','w','E_before','E_after','ell','pv','actual_price','normal_fee','adjustment_fee','emergency_fee','total_fee','contract_version','published_position']

def csvwrite(path,fields,rows):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)

def intervals(day):
    def label(m): return f'{m//60:02d}:{m%60:02d}'
    return [(label(k*240)+'-'+label((k+1)*240),k*24,(k+1)*24) for k in range(6)]

def events(rows):
    result=[]; current=None
    for r in rows:
        i=r['position']
        if r['g']>0:
            if current is None or current['end']!=i or current['day']!=i//144:
                current=dict(day=i//144,start=i,end=i+1,energy=float(r['g'])); result.append(current)
            else: current['end']=i+1; current['energy']+=r['g']
        else: current=None
    return result

def export_book(path,rows):
    by={r['position']:r for r in rows}; orders=[]; storage=[]; emg=[]; provenance=[]
    template=ROOT/'C题/附件/附件5/result4-2.xlsx'
    for d in range(31,365):
        date=(BASE+timedelta(days=d)).strftime('%Y-%m-%d')
        orders.append([date]+[by[i]['q0'] if i in by else None for i in range(d*144+1,d*144+145)]+[math.fsum(by[i][k] for i in range(d*144+1,d*144+145)) if d<364 else None for k in ('q0','normal')])
        for t in range(144):
            i=d*144+t+1
            if i in by: provenance.append(dict(excel_row=d-29,excel_column=t+2,position=i,published=by[i]['published'],field='q0'))
        for k,(label,left,right) in enumerate(intervals(d)):
            storage.append([date,label,math.fsum(by[d*144+t]['c'] for t in range(left,right)),math.fsum(by[d*144+t]['s'] for t in range(left,right)),'00:00' if k==0 else ('24:00' if k==1 else None),by[d*144]['E_before'] if k==0 else (by[d*144+143]['E_after'] if k==1 else None)])
    for e in events(rows):
        d=e['day']; a=(e['start']-d*144)*10; z=(e['end']-d*144)*10
        emg.append([(BASE+timedelta(days=d)).strftime('%Y-%m-%d'),f'{a//60:02d}:{a%60:02d}-{z//60:02d}:{z%60:02d}',e['energy']])
    payload=path.with_suffix('.payload.json'); save(payload,dict(template=str(template),sheets=[orders,storage,emg]))
    node=Path(__import__('shutil').which('node') or 'node')
    subprocess.run([str(node),str(HERE/'workbook.mjs'),str(payload),str(path)],check=True,timeout=90)
    csvwrite(path.parent/'cell_sources.csv',['excel_row','excel_column','position','published','field'],provenance)
    save(path.parent/'boundary_notes.json',dict(status='DRAFT_MISSING_2026_CONTRACT',blank_cells=['EO335','EP335','EQ335'],reason='2026 first interval formal contract and Q42 endpoint actual price absent; dependent window totals blank',natural_day_range=[4464,52560],first_omitted_display_interval=4464,template_sha256=sha(template),planned_fee_column='actual price * normal order; emergency fee separate'))

def readback(path,day_files):
    """Reload contracts and actions independently; do not consume writer provenance."""
    contracts={}; actions={}
    for file in day_files:
        b=unpack(file)
        for v in b['versions']:
            assert v['published_position']==v['start'] and v['end']-v['start']==144
            for k,q in enumerate(v['orders']): contracts[v['start']+k]=float(q)
        for r in b['ledger']: actions[r['position']]=r
    wb=load_workbook(path,data_only=True); order,storage,emergency=wb.worksheets; count=0; worst=0.
    def compare(cell,value):
        nonlocal worst,count
        assert isinstance(cell.value,(int,float)), cell.coordinate
        e=abs(cell.value-value); assert e<1e-6,(cell.coordinate,e); worst=max(worst,e); count+=1
    for row in range(2,336):
        date=order.cell(row,1).value; day=(date-BASE).days
        for col in range(2,146):
            # Derive interval from template column labels, including explicit +1 midnight.
            start=str(order.cell(1,col).value).split('-')[0]; h,m=map(int,start.split(':'))
            i=day*144+h*6+m//10+(144 if col==145 else 0)
            if i>=52560: assert order.cell(row,col).value is None
            else: compare(order.cell(row,col),contracts[i])
        if day<364:
            ids=range(day*144+1,day*144+145)
            compare(order.cell(row,146),math.fsum(contracts[i] for i in ids))
            compare(order.cell(row,147),math.fsum(contracts[i]*actions[i]['price'] for i in ids))
        else: assert order.cell(row,146).value is None and order.cell(row,147).value is None
    for row in range(2,2006):
        day=(storage.cell(row,1).value-BASE).days; text=storage.cell(row,2).value
        a,z=text.split('-'); ah,am=map(int,a.split(':')); zh,zm=map(int,z.split(':')); ids=range(day*144+ah*6+am//10,day*144+zh*6+zm//10)
        compare(storage.cell(row,3),math.fsum(actions[i]['c'] for i in ids)); compare(storage.cell(row,4),math.fsum(actions[i]['s'] for i in ids))
        if storage.cell(row,5).value=='00:00': compare(storage.cell(row,6),actions[day*144]['E_before'])
        elif storage.cell(row,5).value=='24:00': compare(storage.cell(row,6),actions[day*144+143]['E_after'])
        else: assert storage.cell(row,6).value is None
    covered=set(); event_count=0
    for row in range(2,emergency.max_row+1):
        if emergency.cell(row,2).value is None: continue
        day=(emergency.cell(row,1).value-BASE).days; a,z=emergency.cell(row,2).value.split('-'); ah,am=map(int,a.split(':')); zh,zm=map(int,z.split(':'))
        ids=list(range(day*144+ah*6+am//10,day*144+zh*6+zm//10)); assert ids and not covered.intersection(ids)
        assert all(actions[i]['g']>0 for i in ids); covered.update(ids); compare(emergency.cell(row,3),math.fsum(actions[i]['g'] for i in ids)); event_count+=1
    assert covered=={i for i,r in actions.items() if r['g']>0}; wb.close()
    return dict(status='PASS_DRAFT_READBACK',checked_numeric_cells=count,max_error=worst,events=event_count,missing_cells=['EO335','EP335','EQ335'],workbook_sha256=sha(path))

def finalize(out,jan,budget,export_xlsx=False):
    from audit_kernel import raw_sources
    from run import check_day
    raw=raw_sources(); summary=[]
    for branch in BRANCHES:
        budget.check(); dest=out/branch; files=[dest/f'days/{d:03d}.q4z' for d in range(31,365)]; rows=[]; reports=[]; E=10800.; metric={}
        for f in files:
            budget.check(); bb=unpack(f); rep=check_day(bb,raw,E); reports.append(rep); rows.extend(bb['ledger']); E=rep['final_energy']
            for plan in bb['plans']:
                rec=plan['price_record']; r=rec['r']; slot=r%144
                for h in (0,1):
                    pairs=[(int(i),float(p)) for i,p in zip(rec['targets'],rec['prices']) if int(i)//144-r//144==h]
                    if not pairs: continue
                    key=(plan['stage'],h); item=metric.setdefault(key,[0,0.,0.]); errors=[p-raw['price'][i] for i,p in pairs]; item[0]+=len(errors); item[1]+=math.fsum(abs(e) for e in errors); item[2]+=math.fsum(e*e for e in errors)
        assert len(rows)==48096
        standardized=[]
        for r in jan['ledger']+rows:
            i=r['position']; standardized.append(dict(zip(FIELDS,[i,(BASE+timedelta(minutes=10*i)).isoformat(),(BASE+timedelta(minutes=10*(i+1))).isoformat(),r['q0'],r['a'],r['c'],r['s'],r['g'],r['w'],r['E_before'],r['E_after'],r['ell'],r['pv'],r['price'],r['normal'],r['adjustment'],r['emergency'],r['total'],r['version_id'],r['published']])))
        csvwrite(dest/'trajectory.csv',FIELDS,standardized)
        monthly=[]
        for month in range(1,13):
            part=[r for r in jan['ledger']+rows if (BASE+timedelta(minutes=10*r['position'])).month==month]
            monthly.append(dict(month=month,segments=len(part),normal=math.fsum(r['normal'] for r in part),adjustment=0.,emergency=math.fsum(r['emergency'] for r in part),total=math.fsum(r['total'] for r in part),normal_kwh=math.fsum(r['q0'] for r in part),emergency_kwh=math.fsum(r['g'] for r in part),unused_kwh=math.fsum(r['w'] for r in part),initial_energy=part[0]['E_before'],final_energy=part[-1]['E_after']))
        csvwrite(dest/'monthly_summary.csv',list(monthly[0]),monthly)
        csvwrite(dest/'forecast_metrics.csv',['stage','horizon_day','samples','MAE','RMSE'],[dict(stage=k[0],horizon_day=k[1],samples=v[0],MAE=v[1]/v[0],RMSE=math.sqrt(v[2]/v[0])) for k,v in metric.items()])
        report=dict(status='PASS_ANNUAL_INDEPENDENT',segments=len(rows),initial_energy=10800.,final_energy=E,January=jan['total'],Feb_Dec=math.fsum(r['total'] for r in rows),Jan_Dec=jan['total']+math.fsum(r['total'] for r in rows),fallbacks=sum(r['fallbacks'] for r in reports),unproven=sum(r['unproven'] for r in reports),secondary_not_accepted=sum(r['secondary_not_accepted'] for r in reports),numerical_zero_events=sum(r['numerical_zero_events'] for r in reports),max_physical_residual=max(r['max_physical_residual'] for r in reports),max_bill_error=max(r['max_bill_error'] for r in reports))
        save(dest/'independent_audit.json',report)
        if export_xlsx:
            name='result4-2_DRAFT.xlsx' if branch=='M0_P1' else f'{branch}_DRAFT.xlsx'; book=dest/name; export_book(book,rows); save(dest/'export_readback.json',readback(book,files))
        selected=[r for r in standardized if datetime.fromisoformat(r['delivery_start']).strftime('%m-%d') in ('03-20','06-21','09-23','12-21')]
        csvwrite(dest/'four_selected_days.csv',FIELDS,selected); summary.append(dict(branch=branch,**report)); budget.persist()
    save(out/'summary.json',dict(status='PASS_ANNUAL_DRAFT_EXPORT' if export_xlsx else 'PASS_ANNUAL_CSV_AUDIT',branches=summary,comparisons={'R2_P1_minus_M0_P1':summary[1]['Feb_Dec']-summary[0]['Feb_Dec'],'M0_P1_minus_M0_P7':summary[0]['Feb_Dec']-summary[2]['Feb_Dec']},formal_submission_complete=False))
