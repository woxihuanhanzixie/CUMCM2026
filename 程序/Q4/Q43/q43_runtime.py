#!/usr/bin/env python3
"""Q4-3 runtime: forward-hold View, regenerated Bases, P1/P7/P0 Prices, annual Budget.

Adapted from teammate Q4 P1 prereq_tests/runtime.py. Changes vs the verified
small-test version:
  * ROOT/attachments point at this project (no D:/Competition hardcode).
  * Bases.get Q43 branch regenerates load/PV base via base_predict (same frozen
    forecasting.py), instead of reading the teammate annual cache (which shipped
    without its commits/ base-id map).
  * Budget is the annual budget (38000 LP / 2250 fits / 20 min), no cumulative
    small-test counters, no file persistence of coefficients during the run.
"""
from __future__ import annotations

import hashlib
import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
VENDOR = HERE / 'vendor'
sys.path.insert(0, str(VENDOR))

from core import digest  # noqa: E402
from data_access import minutes, Access  # noqa: E402
from forecasting import array_hash, base_predict  # noqa: E402

PROJ = next(p for p in Path(__file__).resolve().parents if (p/'C题'/'附件'/'附件1.xlsx').is_file()) / 'C题'
ATT = PROJ / '附件'


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def sheet(src: Path, name: str) -> np.ndarray:
    wb = load_workbook(src, read_only=True, data_only=True)
    rows = list(wb[name].values)
    wb.close()
    assert [minutes(v) for v in rows[0][1:]] == list(range(10, 1441, 10))
    assert [(r[0].date() - datetime(2025, 1, 1).date()).days for r in rows[1:]] == list(range(365))
    a = np.array([r[1:] for r in rows[1:]], float).ravel()
    assert a.shape == (52560,) and np.isfinite(a).all()
    return a


def load_inputs() -> dict:
    a2 = ATT / '附件2.xlsx'
    a3 = ATT / '附件3.xlsx'
    a4 = ATT / '附件4.xlsx'
    load = sheet(a2, '小区负载')
    pv = sheet(a2, '光伏发电实际功率')
    price = sheet(a4, 'Sheet1')
    assert min(load) >= 0 and min(pv) >= 0 and min(price) > 0
    wb = load_workbook(a3, read_only=True, data_only=True)
    rows = list(wb.active.values)
    wb.close()
    official: dict = {}
    day = None
    for row in rows[1:]:
        if row[0] is not None and str(row[0]).strip():
            day = (datetime.strptime(str(row[0]).split()[0], '%Y-%m-%d')
                   - datetime(2025, 1, 1)).days
        key = (day, minutes(row[1]) // 10)
        a = np.array(row[2:], float)
        assert key not in official and a.shape == (24,) and np.isfinite(a).all() and min(a) >= 0
        official[key] = a
    assert set(official) == {(d, s) for d in range(365) for s in (0, 36, 72, 108)}
    return dict(load_nodes=load, pv_nodes=pv, price_nodes=price, official=official,
                prices=np.ones(144),
                input_sha256={'附件2': sha(a2), '附件3': sha(a3), '附件4': sha(a4)})


class Budget:
    def __init__(self, limits: dict):
        self.limits = limits
        self.start = time.perf_counter()
        self.calls = 0
        self.fits = 0
        self.timings = {}

    def remaining(self) -> float:
        return self.limits['wall_seconds'] - (time.perf_counter() - self.start)

    def check(self) -> None:
        if self.remaining() <= 0:
            raise RuntimeError('wall budget exhausted')

    def reserve(self, kind: str, label: str) -> None:
        self.check()
        kind = 'fit' if kind == 'fit' else 'lp'
        if kind == 'fit':
            if self.fits >= self.limits['fits']:
                raise RuntimeError('fits exhausted')
            self.fits += 1
        else:
            if self.calls >= self.limits['calls']:
                raise RuntimeError('calls exhausted')
            self.calls += 1

    def addtime(self, k: str, s: float) -> None:
        self.timings[k] = self.timings.get(k, 0) + s

    def record(self) -> dict:
        return dict(calls=self.calls, fits=self.fits,
                    wall_seconds=time.perf_counter() - self.start, timings=self.timings)


class View(Access):
    """Q43 forward-hold view: completed actuals use node j = max(i-1, 0)."""

    def __init__(self, data, policy='C2', position=0):
        super().__init__(data, policy, position)
        self.mapping = 'Q43'
        self.__raw = {v: np.asarray(data[k]).copy() for v, k in
                      (('load', 'load_nodes'), ('pv', 'pv_nodes'), ('price', 'price_nodes'))}

    def completed(self, end, variable):
        if not 0 <= end <= self.position:
            raise ValueError('future completed history denied')
        ids = np.maximum(np.arange(end) - 1, 0)
        return self.__raw[variable][ids].copy()

    def observations_completed_by(self, end, variable):
        a = self.completed(end, variable)
        self._event('completed', end=end, variable=variable)
        return a

    def history_before(self, day, variable):
        return self.completed(day * 144, variable).reshape(day, 144)

    def delivered(self):
        if self.committed is None:
            raise ValueError('actual settlement before commitment')
        j = max(self.position - 1, 0)
        self._event('delivery', source=j)
        return tuple(float(self.__raw[k][j]) for k in ('load', 'pv', 'price'))


class Bases:
    """Regenerate causal load/PV base per publish day (k=0 today, k=1 tomorrow)."""

    def __init__(self, budget):
        self.budget = budget
        self.hot: dict = {}
        self.refs: dict = {}

    def get(self, view, day):
        if view.position < day * 144:
            raise ValueError('future base publication')
        key = (view.mapping, day)
        hist = {v: view.history_before(day, v) for v in ('load', 'pv')}
        hh = {v: array_hash(x) for v, x in hist.items()}
        if key in self.hot:
            if self.refs[key]['histories'] != hh:
                raise ValueError('cache completed history mismatch')
            return [x.copy() for x in self.hot[key]]
        if day == 1:
            vals = [np.tile(hist[v][-1], 2) for v in ('load', 'pv')]
            meta = {'histories': hh, 'method': 'short-history previous day', 'day': day}
        else:
            vals = []
            for v in ('load', 'pv'):
                preds = []
                for k in range(1 if day == 364 else 2):
                    y, _rec = base_predict(hist[v], day, k, v == 'load', self.budget)
                    preds.append(y)
                vals.append(np.concatenate(preds))
            meta = {'histories': hh, 'method': 'base_predict regenerated', 'day': day}
        want = 144 if day == 364 else 288
        assert all(x.shape == (want,) and np.isfinite(x).all() and min(x) >= 0 for x in vals)
        self.hot[key] = vals
        self.refs[key] = meta
        return [x.copy() for x in vals]


class Prices:
    """P1 weekly-difference OLS + AR(1); P7 last-week same-slot; P0 short-history mean."""

    def __init__(self, budget):
        self.budget = budget
        self.coeffs: dict = {}
        self.records: list = []

    @staticmethod
    def validate_coefficient(c):
        if not all(k in c and np.isfinite(c[k]) for k in ('a', 'b', 'rho', 'rho_raw')):
            raise ValueError('invalid price coefficient')
        if not 0 <= c['rho'] <= 1:
            raise ValueError('invalid price coefficient rho')

    def coefficient(self, view, day):
        r0 = day * 144
        p = view.completed(r0, 'price')
        n = (view.completed(r0, 'load') - view.completed(r0, 'pv')) / 1000
        if not np.isfinite(p).all() or not np.isfinite(n).all():
            raise ValueError('invalid actual history')
        ids = np.arange(max(1008, r0 - 4032), r0)
        if len(ids) < 1008:
            return None
        key = f'{view.mapping}:{day}'
        history = digest([p, n])
        if key in self.coeffs:
            c = self.coeffs[key]
            self.validate_coefficient(c)
            if (c['history'] != history or c['day'] != day or c['last'] != r0 - 1
                    or c['first'] != max(1008, r0 - 4032)):
                raise ValueError('price coefficient history identity mismatch')
            return c
        x = n[ids] - n[ids - 1008]
        y = p[ids] - p[ids - 1008]
        dx = x - x.mean()
        den = dx @ dx
        self.budget.reserve('fit', key + '_OLS')
        b = float(dx @ (y - y.mean()) / den) if den > 1e-12 else 0.
        a = float(y.mean() - b * x.mean())
        e = y - a - b * x
        self.budget.reserve('fit', key + '_AR')
        denrho = e[:-1] @ e[:-1]
        raw = float(e[:-1] @ e[1:] / denrho) if denrho > 1e-12 else 0.
        c = dict(a=a, b=b, rho=float(np.clip(raw, 0, 1)), rho_raw=raw,
                 history=history, day=day, first=int(ids[0]), last=int(ids[-1]),
                 rows=len(ids), degenerate_b=bool(den <= 1e-12),
                 degenerate_rho=bool(denrho <= 1e-12))
        self.coeffs[key] = c
        return c

    def predict(self, view, targets, nhat, method, source):
        r = view.position
        targets = np.asarray(targets, int)
        nhat = np.asarray(nhat, float)
        if method not in ('P1', 'P7') or len(targets) != len(nhat) or not np.isfinite(nhat).all():
            raise ValueError('invalid forecast')
        if np.any(targets < r) or np.any(targets >= min((r // 144 + 2) * 144, 52560)):
            raise ValueError('forecast horizon outside law')
        p = view.completed(r, 'price')
        n = (view.completed(r, 'load') - view.completed(r, 'pv')) / 1000
        if not np.isfinite(p).all() or not np.isfinite(n).all() or np.any(p <= 0):
            raise ValueError('invalid completed price/net history')
        c = self.coefficient(view, r // 144) if method == 'P1' else None
        vals, modes, src, floor = [], [], [], 0
        for i, N in zip(targets, nhat):
            lag = int(i - 1008)
            if c is not None:
                if not 0 <= lag < r:
                    raise ValueError('illegal week source')
                residual = p[-1] - p[-1009] - c['a'] - c['b'] * (n[-1] - n[-1009])
                v = p[lag] + c['a'] + c['b'] * (N - n[lag]) + c['rho'] ** int(i - r + 1) * residual
                floor += v < 1e-6
                vals.append(max(1e-6, float(v)))
                modes.append('P1')
                src.append([lag, r - 1, r - 1009])
            elif 0 <= lag < r:
                vals.append(float(p[lag]))
                modes.append('P7')
                src.append([lag])
            else:
                idx = list(range(int(i % 144), r, 144))[-7:] or list(range(max(0, r - 144), r))
                if not idx:
                    raise ValueError('no price history: do not optimize Jan1 midnight')
                vals.append(float(np.mean(p[idx])))
                modes.append('P0')
                src.append(idx)
        record = dict(mapping=view.mapping, r=r, targets=targets, nhat_mw=nhat,
                      source=source, method=method, coefficient=c, modes=modes,
                      actual_sources=src, floor_count=int(floor), prices=np.array(vals),
                      history_last=r - 1)
        self.records.append(record)
        return np.array(vals), record
