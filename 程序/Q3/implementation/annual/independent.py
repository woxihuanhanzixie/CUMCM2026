"""Independent raw-cell/day stream audit; does not call control or mapping writers."""
import io, math, hashlib
from collections import defaultdict
from datetime import datetime,timedelta
import numpy as np
from openpyxl import load_workbook
from data_access import shared_bytes
from audit import physical
from core import digest
from runtime import read

def raw_sources(root,sources):
    blobs={}
    for rel,h in sources.items():
        b=shared_bytes(root/rel)
        if hashlib.sha256(b).hexdigest()!=h: raise ValueError('Independent source identity failure')
        blobs[rel]=b
    wb=load_workbook(io.BytesIO(blobs['C题/附件/附件2.xlsx']),read_only=True,data_only=True); raw={}
    for sheet,var in [('小区负载','ell'),('光伏发电实际功率','pv')]:
        rows=list(wb[sheet].values)
        if len(rows)!=366: raise AssertionError('Raw date count')
        for day,row in enumerate(rows[1:]):
            if row[0].date()!=(datetime(2025,1,1)+timedelta(days=day)).date(): raise AssertionError('Raw date identity')
        raw[var+'_power']=np.array([float(x) for row in rows[1:] for x in row[1:145]])
        raw[var]=raw[var+'_power']/6
    wb.close(); wb=load_workbook(io.BytesIO(blobs['C题/附件/附件1.xlsx']),read_only=True,data_only=True)
    rows=list(wb.active.values); col=list(rows[0]).index('电价'); raw['prices']={}
    for row in rows[1:]:
        t=str(row[0]); hour,minute=map(int,t.replace('+1','').split(':')[:2]); raw['prices'][(hour%24)*60+minute]=float(row[col])
    wb.close(); wb=load_workbook(io.BytesIO(blobs['C题/附件/附件3.xlsx']),read_only=True,data_only=True)
    raw['official']={}; day=None
    for row in list(wb.active.values)[1:]:
        if row[0] is not None and str(row[0]).strip():
            year,month,date=map(int,str(row[0]).split()[0].split('-'))
            day=(datetime(year,month,date)-datetime(2025,1,1)).days
        hour=int(str(row[1]).split(':')[0]); raw['official'][day,hour*6]=np.array(row[2:],dtype=float)
    wb.close(); return raw

def close(a,b,tol=1e-6):
    if np.shape(a)!=np.shape(b) or not np.isfinite(np.asarray(a,dtype=float)).all() or not np.isfinite(np.asarray(b,dtype=float)).all() or np.max(np.abs(np.asarray(a)-np.asarray(b)),initial=0)>tol: raise AssertionError('Numeric mismatch')

def audit_base_values(base,raw):
    """Rebuild features/normal-equation residual from raw cells, without fitting."""
    key=base['key']; fingerprint=digest(base)
    checked=raw.setdefault('_checked_bases',{})
    if checked.get(key)==fingerprint: return
    d=base['day']
    for task in base['tasks']:
        record=task['record']; k=task['lead']; load=task['variable']=='load'; var='ell' if load else 'pv'
        history=raw[var+'_power'][np.maximum(np.arange(d*144)-1,0)].reshape(d,144)
        def features(h):
            slots=np.arange(144); mean3=np.mean(history[h-3:h],axis=0)
            cols=[history[h-1],history[h+k-7],history[h+k-14],mean3,np.mean(history[h-7:h],axis=0),mean3-np.mean(history[h-6:h-3],axis=0),np.repeat(np.mean(history[h-1]),144)]
            for harmonic in (1,2,3): cols.extend([np.sin(2*np.pi*harmonic*slots/144),np.cos(2*np.pi*harmonic*slots/144)])
            if load:
                weekday=(datetime(2025,1,1)+timedelta(days=h+k)).weekday()
                cols.extend([np.repeat(float(weekday==w),144) for w in range(1,7)])
            return np.stack(cols,axis=1)
        hs=list(range(max(14,d-28-k),d-k))
        if len(hs)<7:
            src=d+k-7 if load and 0<=d+k-7<d else d-1
            if record['fallback_day']!=src: raise AssertionError('Fallback history day')
            close(record['raw'],history[src],0)
        else:
            X=np.concatenate([features(h) for h in hs]); y=np.concatenate([history[h+k] for h in hs]); w=np.repeat([2**(-(d-1-h-k)/14) for h in hs],144)
            mu=np.average(X,axis=0,weights=w); sd=np.sqrt(np.average((X-mu)**2,axis=0,weights=w)); scale=np.where(sd<1e-12 if load else sd<=1e-12,1,sd)
            close(mu,record['mean']); close(scale,record['scale']); beta=np.array(record['beta']); intercept=record['intercept']; Z=(X-mu)/scale; err=intercept+Z@beta-y
            residual=max(abs(w@err),np.max(abs(Z.T@(w*err)+beta)))/max(1,np.max(abs(np.c_[np.ones(len(y)),Z].T@(w*y))))
            if residual>1e-7: raise AssertionError('Independent ridge normal-equation residual')
            close(record['raw'],intercept+(features(d)-mu)/scale@beta)
    # Keep only the latest day: the independent stream is bounded in memory too.
    raw['_checked_bases']={key:fingerprint}

def audit_block(block,raw,initial_energy,start=None,end=None,cache_dir=None):
    rows=block['ledger']; policy=block['policy']
    if not rows: raise AssertionError('Empty block')
    start=rows[0]['position'] if start is None else start; end=rows[-1]['position']+1 if end is None else end
    if [r['position'] for r in rows]!=list(range(start,end)): raise AssertionError('Missing/duplicate deliveries')
    if start//144!=(end-1)//144: raise AssertionError('Day stream crosses midnight')
    day=start//144; plans={p['position']:p for p in block['plans']}
    expected=[i for i in range(start,end) if i%6==0 and i>=36]
    if list(plans)!=expected or len(plans)!=len(block['plans']): raise AssertionError('Wrong LP count/clock')
    versions=block['versions']; day_versions=[v for v in versions if v['day']==day]
    required=[day*144+s for s in ((0,36,72,108) if policy=='C2' else (0,)) if day*144+s<end]
    if [v['published_position'] for v in day_versions]!=required: raise AssertionError('Official contract publication sequence')
    original=day_versions[0]; q=np.array(original['orders'])
    for v in day_versions:
        s=v['slot']
        if v['policy']!=policy or v['start']!=day*144+s or v['end']!=(day+1)*144 or v['published_position']!=v['start'] or len(v['orders'])!=144-s or v['kind']!=('original' if s==0 else 'adjust'): raise AssertionError('Invalid contract identity')
        close(v['q0'],q,0)
    log=defaultdict(list)
    for e in block['access_log']:
        i=e['position']; log[i].append(e)
        if e['kind'].startswith('DENIED'): raise AssertionError('Denied access in real run')
        if e['kind']=='history' and (e['day']!=day or day*144>i): raise AssertionError('History cutoff')
        if e['kind']=='completed' and e['end']>i: raise AssertionError('Future access')
        if e['kind']=='official':
            if e['day']!=day or e['slot'] not in (0,36,72,108) or i!=day*144+e['slot'] or (policy=='C0' and e['slot']!=0): raise AssertionError('Official permission')
            h=hashlib.sha256(np.asarray(raw['official'][day,e['slot']],dtype='<f8').tobytes()).hexdigest()
            if h!=e['sha256']: raise AssertionError('Official source hash')
        if e['kind'] in ('history','completed'):
            n=e['day']*144 if e['kind']=='history' else e['end']; var='ell' if e['variable']=='load' else 'pv'
            if n>i: raise AssertionError('History beyond completed interval')
            powers=raw[var+'_power'][np.maximum(np.arange(n)-1,0)]
            if hashlib.sha256(np.asarray(powers,dtype='<f8').tobytes()).hexdigest()!=e['sha256']: raise AssertionError('Historical observations changed')
    if start==day*144:
        releases=[e['slot'] for e in block['access_log'] if e['kind']=='official']
        if releases!=([0] if policy=='C0' else [0,36,72,108]): raise AssertionError('Missing/duplicate official release')
    forecasts=block['forecasts']; fmap={f['position']:f for f in forecasts}
    base=None
    if cache_dir is not None and day>0:
        keys={r['base_id'] for r in rows}
        if len(keys)!=1: raise AssertionError('Midnight base changed')
        key=keys.pop(); base=read(cache_dir/(key+'.json.gz'))
        if base['day']!=day or base['history_cutoff']!=day*144 or base['key']!=key: raise AssertionError('Base cutoff identity')
        audit_base_values(base,raw)
        for variable,var in [('load','ell'),('pv','pv')]:
            hist=raw[var+'_power'][np.maximum(np.arange(day*144)-1,0)]
            if hashlib.sha256(np.asarray(hist,dtype='<f8').tobytes()).hexdigest()!=base['history_hashes'][variable]: raise AssertionError('Cache history differs from raw source')
        for t in base['tasks']:
            rr=t['record']; k=t['lead']; hs=list(range(max(14,day-28-k),day-k))
            if rr['training_days']!=hs or rr['target_days']!=[h+k for h in hs] or t['target_day']!=day+k or t['history_cutoff']!=day*144: raise AssertionError('Cache training identity')
            if rr['method']!=('ridge' if len(hs)>=7 else 'fallback') or rr.get('stationarity',0)>1e-7: raise AssertionError('Cache fit certificate')
            close(np.maximum(rr['raw'],0),base['predictions'][t['variable']][k*144:(k+1)*144])
        events=[e for e in block['access_log'] if e['kind']=='base_cache']
        if start==day*144 and (len(events)!=1 or events[0]['key']!=key or events[0]['history_cutoff']!=day*144): raise AssertionError('Missing cache read identity')
        for e in events:
            if e['cache_sha256']!=hashlib.sha256((cache_dir/(key+'.json.gz')).read_bytes()).hexdigest(): raise AssertionError('Cache file changed after use')
        # Independent scalar fusion checks all saved points against raw bulletins.
        for f in forecasts:
            rslot=0 if policy=='C0' else f['position']%144
            bulletin=raw['official'][day,rslot]; anchor=raw['pv_power'][day*144+rslot-2]
            baseline=base['predictions']['pv']; curve=list(baseline); nodes=[anchor]+list(bulletin)
            for m in range(rslot,min(len(curve),rslot+145)):
                lead=m-rslot; h,part=divmod(lead,6)
                if part==0: curve[m]=nodes[h]
                else:
                    weight=part/6; left=rslot+6*h; right=min(left+6,len(baseline)-1)
                    curve[m]=max(0,baseline[m]+(1-weight)*(nodes[h]-baseline[left])+weight*(nodes[h+1]-baseline[right]))
            close(f['pv'],curve,1e-6); close(f['load'],base['predictions']['load'])
    energy=initial_energy; maxres=0.; fields=('normal','adjustment','emergency','total','c','s','g','w'); totals={k:[] for k in fields}
    for r in rows:
        i=r['position']; slot=i%144; stamp=datetime(2025,1,1)+timedelta(minutes=i*10); node=0 if i==0 else i-1
        if r['policy']!=policy or r['day']!=day or r['slot']!=slot or r['timezone']!='Asia/Shanghai' or r['delivery_start']!=stamp.isoformat() or r['delivery_end']!=(stamp+timedelta(minutes=10)).isoformat() or r['source_node_index']!=node or r['initial_missing_node_approximation']!=(i==0): raise AssertionError('Time/source mapping')
        if r['display_row_date']!=(stamp-timedelta(minutes=10)).date().isoformat() or r['q0_created_at']!=day*144 or r['order_owner_day']!=day: raise AssertionError('Contract/display owner')
        L,V,price=raw['ell'][node],raw['pv'][node],raw['prices'][slot*10]
        close([r['ell'],r['pv'],r['price'],r['E_before']],[L,V,price,energy])
        v=[v for v in day_versions if v['start']<=i][-1]; A=v['orders'][i-v['start']]
        close([r['a_final'],r['q0']],[A,q[slot]])
        if r['final_version_id']!=v['id'] or r['final_published_at']!=v['published_position']: raise AssertionError('Last valid contract')
        if policy!='C2': close(A,q[slot],0)
        events=log[i]; kinds=[e['kind'] for e in events]
        if kinds.count('commit')!=1 or kinds.count('reveal')!=1 or kinds.index('commit')>kinds.index('reveal'): raise AssertionError('Commit/reveal order')
        close(events[kinds.index('commit')]['order'],A)
        if events[kinds.index('reveal')]['source_node_index']!=node: raise AssertionError('Reveal source')
        release=0 if policy=='C0' else (slot//36)*36
        if list(r['official_version'])!=[day,release] or r['history_cutoff']>i: raise AssertionError('Wrong forecast identity')
        R=1200 if i<36 else plans[i-i%6]['solution']['E'][i%6]
        close(R,r['R'])
        if i>=36 and r['hour_plan_time']!=i-i%6: raise AssertionError('Reference clock')
        # Independently verify feedback rule in addition to feasibility.
        surplus=A+V-L
        if surplus>=0: c=min(surplus,5000/6,(10800-energy)/.9); s=g=0.; w=surplus-c
        else: c=w=0.; s=min(-surplus,5000/6,.9*max(energy-R,0)); g=-surplus-s
        close([r[k] for k in ('c','s','g','w')],[c,s,g,w])
        maxres=max(maxres,physical(energy,L,V,A,dict(c=c,s=s,g=g,w=w,E=r['E_after'])))
        energy=r['E_after']; fees=[price*A,.5*price*abs(A-q[slot]),5*price*g]; total=math.fsum(fees)
        close([r[k] for k in fields[:4]],fees+[total])
        for k,value in zip(fields,fees+[total,c,s,g,w]): totals[k].append(value)
    for i,plan in plans.items():
        p=plan['problem']; sol=plan['solution']; slot=i%144; n=min((day+2)*144,52560)-i; stage='original' if slot==0 else ('adjust' if policy=='C2' and slot in (36,72,108) else 'feedback')
        if plan['stage']!=stage or p['stage']!=stage or p['today']!=144-slot or len(p['ell'])!=n or len(sol['E'])!=n or plan['internal_b_committed']: raise AssertionError('LP range/stage')
        if sol['input_hash']!=digest(p) or abs(sol['upper']-sol['lower'])>1e-5 or sol['residual']>1e-6 or sol['stationarity']>1e-7: raise AssertionError('LP certificate')
        close(p['energy'],rows[i-start]['E_before']); close(p['prices'],[raw['prices'][(slot+j)%144*10] for j in range(n)])
        if stage!='original':
            close(p['q0'],q[slot:]); before=[v for v in day_versions if v['start']<i or v['slot']==0][-1]
            close(p['current'],before['orders'][i-before['start']:])
        if stage=='feedback': close(sol['A'][:144-slot],p['current'])
        else: close(sol['A'][:144-slot],[v for v in day_versions if v['start']==i][0]['orders'])
        available=[f for f in forecasts if f['position']<=i]
        if not available: raise AssertionError('No frozen forecast')
        forecast=available[-1]; close(p['ell'],np.array(forecast['load'][slot:slot+n])/6); close(p['pv'],np.array(forecast['pv'][slot:slot+n])/6)
        if list(plan['official_version'])!=list(forecast['official_version']) or plan['base_id']!=forecast.get('base_id') or plan['history_cutoff']!=forecast['history_cutoff']: raise AssertionError('Plan forecast link')
        En=p['energy']; fees=[]
        for t in range(n):
            act={k:sol[k][t] for k in ('c','s','g','w')}; act['E']=sol['E'][t]
            physical(En,p['ell'][t],p['pv'][t],sol['A'][t],act); En=sol['E'][t]
            A=sol['A'][t]; U=p['ell'][t]+5000/6
            if stage!='original' and t<p['today']: U=max(U,p['q0'][t],p['current'][t])
            if A>U+1e-6: raise AssertionError('LP order upper bound')
            penalty=.5*abs(A-p['q0'][t]) if stage!='original' and t<p['today'] else 0
            fees.append(p['prices'][t]*(A+penalty+5*sol['g'][t]))
        close(math.fsum(fees),sol['upper'])
    sums={k:math.fsum(v) for k,v in totals.items()}
    close(sums['total'],math.fsum(r['total'] for r in rows),1e-4)
    return dict(status='PASS',policy=policy,day=day,segments=end-start,plans=len(plans),initial_SOC=initial_energy,final_SOC=energy,max_physical_residual=maxres,**sums)
