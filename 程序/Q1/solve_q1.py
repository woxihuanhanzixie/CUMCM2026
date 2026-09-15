#!/usr/bin/env python3
"""Recovered Q1 solver: LP, independent LP, MILP fallback and exports."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from datetime import time as dt_time
from pathlib import Path
from typing import Any

import numpy as np
import scipy
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, linprog, milp

T = 144
DT = 1.0 / 6.0
ETA = 0.9
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_INITIAL = 6000.0
POWER_MAX = 5000.0
ENERGY_MAX = POWER_MAX * DT
TOL = 1e-7
COST_TOL = 1e-6
PROJECT_ROOT = next(p for p in Path(__file__).resolve().parents if (p/'C题'/'附件'/'附件1.xlsx').is_file())
Q1_ROOT = Path(__file__).resolve().parent
DEFAULT_ATTACHMENT = PROJECT_ROOT / "C题" / "附件" / "附件1.xlsx"
DEFAULT_TEMPLATE = PROJECT_ROOT / "C题" / "附件" / "附件5" / "result1.xlsx"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def point_minutes(value: Any) -> int:
    if isinstance(value, dt_time):
        return value.hour * 60 + value.minute
    text = str(value).strip().replace("：", ":")
    plus_one = text.endswith("+1")
    if plus_one:
        text = text[:-2]
    hour, minute = (int(part) for part in text.split(":", 1))
    total = hour * 60 + minute + (24 * 60 if plus_one else 0)
    if total < 0 or total > 24 * 60 or total % 10:
        raise ValueError(f"invalid time label: {value!r}")
    return total


def format_minutes(minutes: int) -> str:
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def canonical_interval(t: int) -> tuple[str, str]:
    start, end = (t - 1) * 10, t * 10
    return format_minutes(start), ("00:00" if end == 1440 else format_minutes(end))


def load_attachment(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ValueError(f"unexpected sheets: {workbook.sheetnames}")
        worksheet = workbook["Sheet1"]
        if worksheet.max_row != 145 or worksheet.max_column < 4:
            raise ValueError(f"unexpected dimensions: {worksheet.max_row}x{worksheet.max_column}")
        rows = list(worksheet.iter_rows(min_row=2, max_row=145, values_only=True))
    finally:
        workbook.close()
    price, load_power, pv_power = [], [], []
    for index, row in enumerate(rows, start=1):
        if point_minutes(row[0]) != index * 10:
            raise ValueError(f"time order error at t={index}")
        p, load, pv = float(row[1]), float(row[2]), float(row[3])
        if not all(math.isfinite(value) for value in (p, load, pv)):
            raise ValueError(f"nonfinite value at t={index}")
        if p <= 0 or load < 0 or pv < 0:
            raise ValueError(f"invalid value at t={index}")
        price.append(p); load_power.append(load); pv_power.append(pv)
    return {
        "price": price,
        "load_power": load_power,
        "pv_power": pv_power,
        "load_energy": [value * DT for value in load_power],
        "pv_energy": [value * DT for value in pv_power],
    }

def solve_lp(data: dict[str, Any]) -> dict[str, Any]:
    prices = np.asarray(data["price"], float)
    load = np.asarray(data["load_energy"], float)
    pv = np.asarray(data["pv_energy"], float)
    nq = nc = ns = nw = T
    ne = T + 1
    size = nq + nc + ns + nw + ne
    q0, c0, s0, w0, e0 = 0, nq, nq + nc, nq + nc + ns, nq + nc + ns + nw
    objective = np.zeros(size); objective[q0:q0 + nq] = prices
    Aeq = np.zeros((2 * T, size)); beq = np.r_[load - pv, np.zeros(T)]
    for t in range(T):
        Aeq[t, q0 + t] = 1; Aeq[t, c0 + t] = -1; Aeq[t, s0 + t] = 1; Aeq[t, w0 + t] = -1
        row = T + t
        Aeq[row, e0 + t] = -1; Aeq[row, e0 + t + 1] = 1
        Aeq[row, c0 + t] = -ETA; Aeq[row, s0 + t] = 1 / ETA
    lower = np.zeros(size); upper = np.full(size, np.inf)
    upper[c0:c0 + nc] = ENERGY_MAX; upper[s0:s0 + ns] = ENERGY_MAX; upper[w0:w0 + nw] = pv
    lower[e0:e0 + ne] = SOC_MIN; upper[e0:e0 + ne] = SOC_MAX
    lower[e0] = upper[e0] = SOC_INITIAL; lower[e0 + T] = upper[e0 + T] = SOC_INITIAL
    result = linprog(objective, A_eq=Aeq, b_eq=beq, bounds=list(zip(lower, upper)), method="highs")
    if not result.success:
        raise RuntimeError(f"LP failed: {result.message}")
    q, c, s, w, E = (result.x[q0:q0+nq], result.x[c0:c0+nc], result.x[s0:s0+ns], result.x[w0:w0+nw], result.x[e0:e0+ne])
    diagnostics = {
        "balance_marginal": result.eqlin.marginals[:T].copy(),
        "state_marginal": result.eqlin.marginals[T:].copy(),
        "soc_lower_marginal": result.lower.marginals[e0:e0+ne].copy(),
        "soc_upper_marginal": result.upper.marginals[e0:e0+ne].copy(),
        "charge_upper_marginal": result.upper.marginals[c0:c0+nc].copy(),
        "discharge_upper_marginal": result.upper.marginals[s0:s0+ns].copy(),
    }
    return {"method": "LP", "objective": float(result.fun), "q": q, "c": c, "s": s, "w": w, "E": E, "diagnostics": diagnostics}


def solve_independent(data: dict[str, Any]) -> dict[str, Any]:
    prices = np.asarray(data["price"], float); load = np.asarray(data["load_energy"], float); pv = np.asarray(data["pv_energy"], float)
    nq = nc = ns = nw = T; size = 4 * T
    q0, c0, s0, w0 = 0, nq, nq + nc, nq + nc + ns
    objective = np.zeros(size); objective[q0:q0 + nq] = prices
    Aeq = np.zeros((T + 1, size)); beq = np.zeros(T + 1)
    for t in range(T):
        Aeq[t, q0 + t] = 1; Aeq[t, c0 + t] = -1; Aeq[t, s0 + t] = 1; Aeq[t, w0 + t] = -1
        beq[t] = load[t] - pv[t]
    for i in range(T):
        Aeq[T, c0 + i] = ETA; Aeq[T, s0 + i] = -1 / ETA
    Aub = np.zeros((2 * T, size)); bub = np.r_[np.full(T, SOC_INITIAL - SOC_MIN), np.full(T, SOC_MAX - SOC_INITIAL)]
    for t in range(T):
        for i in range(t + 1):
            Aub[t, c0 + i] = -ETA; Aub[t, s0 + i] = 1 / ETA
            Aub[T + t, c0 + i] = ETA; Aub[T + t, s0 + i] = -1 / ETA
    lower = np.zeros(size); upper = np.full(size, np.inf)
    upper[c0:c0 + nc] = ENERGY_MAX; upper[s0:s0 + ns] = ENERGY_MAX; upper[w0:w0 + nw] = pv
    result = linprog(objective, A_eq=Aeq, b_eq=beq, A_ub=Aub, b_ub=bub, bounds=list(zip(lower, upper)), method="highs")
    if not result.success:
        raise RuntimeError(f"independent LP failed: {result.message}")
    q, c, s, w = result.x[q0:q0+nq], result.x[c0:c0+nc], result.x[s0:s0+ns], result.x[w0:w0+nw]
    E = np.empty(T + 1); E[0] = SOC_INITIAL
    for t in range(T): E[t + 1] = E[t] + ETA * c[t] - s[t] / ETA
    return {"method": "independent-reduced-LP", "objective": float(result.fun), "q": q, "c": c, "s": s, "w": w, "E": E}

def solve_milp(data: dict[str, Any]) -> dict[str, Any]:
    prices = np.asarray(data["price"], float); load = np.asarray(data["load_energy"], float); pv = np.asarray(data["pv_energy"], float)
    nq = nc = ns = nw = nz = T; ne = T + 1; size = 4 * T + ne + nz
    q0, c0, s0, w0, e0, z0 = 0, nq, nq + nc, nq + nc + ns, nq + nc + ns + nw, nq + nc + ns + nw + ne
    objective = np.zeros(size); objective[q0:q0 + nq] = prices
    Aeq = np.zeros((2 * T, size)); beq = np.r_[load - pv, np.zeros(T)]
    for t in range(T):
        Aeq[t, q0 + t] = 1; Aeq[t, c0 + t] = -1; Aeq[t, s0 + t] = 1; Aeq[t, w0 + t] = -1
        row = T + t
        Aeq[row, e0 + t] = -1; Aeq[row, e0 + t + 1] = 1; Aeq[row, c0 + t] = -ETA; Aeq[row, s0 + t] = 1 / ETA
    Aineq = np.zeros((2 * T, size)); lower_ineq = np.full(2 * T, -np.inf); upper_ineq = np.r_[np.zeros(T), np.full(T, ENERGY_MAX)]
    for t in range(T):
        Aineq[t, c0 + t] = 1; Aineq[t, z0 + t] = -ENERGY_MAX
        Aineq[T + t, s0 + t] = 1; Aineq[T + t, z0 + t] = ENERGY_MAX
    lower = np.zeros(size); upper = np.full(size, np.inf)
    upper[c0:c0 + nc] = ENERGY_MAX; upper[s0:s0 + ns] = ENERGY_MAX; upper[w0:w0 + nw] = pv
    lower[e0:e0 + ne] = SOC_MIN; upper[e0:e0 + ne] = SOC_MAX
    lower[e0] = upper[e0] = SOC_INITIAL; lower[e0 + T] = upper[e0 + T] = SOC_INITIAL; upper[z0:z0 + nz] = 1
    integrality = np.zeros(size, dtype=int); integrality[z0:z0 + nz] = 1
    constraints = [LinearConstraint(Aeq, beq, beq), LinearConstraint(Aineq, lower_ineq, upper_ineq)]
    result = milp(objective, integrality=integrality, bounds=Bounds(lower, upper), constraints=constraints)
    if not result.success:
        raise RuntimeError(f"MILP failed: {result.message}")
    return {"method": "MILP", "objective": float(result.fun), "q": result.x[q0:q0+nq], "c": result.x[c0:c0+nc], "s": result.x[s0:s0+ns], "w": result.x[w0:w0+nw], "E": result.x[e0:e0+ne]}


def calculate_checks(result: dict[str, Any], data: dict[str, Any]) -> dict[str, float]:
    q, c, s, w, E = result["q"], result["c"], result["s"], result["w"], result["E"]
    prices = np.asarray(data["price"], float); load = np.asarray(data["load_energy"], float); pv = np.asarray(data["pv_energy"], float)
    return {
        "objective": float(np.dot(prices, q)),
        "balance_residual_max": float(np.max(np.abs(q + pv + s - load - c - w))),
        "state_residual_max": float(np.max(np.abs(E[1:] - E[:-1] - ETA * c + s / ETA))),
        "soc_bound_violation": float(max(0.0, SOC_MIN - np.min(E), np.max(E) - SOC_MAX)),
        "charge_power_violation": float(max(0.0, np.max(c / DT) - POWER_MAX)),
        "discharge_power_violation": float(max(0.0, np.max(s / DT) - POWER_MAX)),
        "nonnegative_violation": float(max(0.0, -np.min(q), -np.min(c), -np.min(s), -np.min(w))),
        "curtailment_violation": float(max(0.0, np.max(w - pv))),
        "mutual_exclusion_violation": float(np.max(c * s)),
        "terminal_soc_error": float(max(abs(E[0] - SOC_INITIAL), abs(E[-1] - SOC_INITIAL))),
        "total_purchase": float(np.sum(q)), "total_charge": float(np.sum(c)),
        "total_discharge": float(np.sum(s)), "total_curtailment": float(np.sum(w)),
        "max_charge_power": float(np.max(c / DT)), "max_discharge_power": float(np.max(s / DT)),
    }

def export_trace(path: Path, result: dict[str, Any], data: dict[str, Any]) -> None:
    columns = ["t","source_timestamp","canonical_start","canonical_end","price","load_power","pv_power","load_energy","pv_energy","q","c","s","w","E_start","E_end","charge_power","discharge_power","purchase_cost"]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns); writer.writeheader()
        for t in range(1, T + 1):
            i = t - 1; start, end = canonical_interval(t)
            row = {"t":t,"source_timestamp":("0:00+1" if t == T else end),"canonical_start":start,"canonical_end":end,
                   "price":data["price"][i],"load_power":data["load_power"][i],"pv_power":data["pv_power"][i],
                   "load_energy":data["load_energy"][i],"pv_energy":data["pv_energy"][i],"q":result["q"][i],
                   "c":result["c"][i],"s":result["s"][i],"w":result["w"][i],"E_start":result["E"][i],"E_end":result["E"][i+1],
                   "charge_power":result["c"][i]/DT,"discharge_power":result["s"][i]/DT,"purchase_cost":data["price"][i]*result["q"][i]}
            writer.writerow({k:(f"{v:.17g}" if isinstance(v,float) else v) for k,v in row.items()})


def export_result(path: Path, template: Path, result: dict[str, Any]) -> None:
    workbook = load_workbook(template)
    try:
        if workbook.sheetnames != ["计划购电量","充放电量"]: raise ValueError("unexpected result sheets")
        plan = workbook["计划购电量"]
        for t in range(1, T + 1):
            start,end=canonical_interval(t); plan.cell(t+1,1).value=f"{start}-{end}"; plan.cell(t+1,2).value=float(result["q"][t-1])
        storage=workbook["充放电量"]
        for block in range(6):
            row=block+2
            storage.cell(row,2).value=float(np.sum(result["c"][block*24:(block+1)*24]))
            storage.cell(row,3).value=float(np.sum(result["s"][block*24:(block+1)*24]))
        storage.cell(2,4).value="0:00"; storage.cell(2,5).value=SOC_INITIAL
        storage.cell(3,4).value="24:00"; storage.cell(3,5).value=SOC_INITIAL
        workbook.save(path)
    finally: workbook.close()


def build_parser():
    parser=argparse.ArgumentParser(description="Solve Q1")
    parser.add_argument("--attachment",type=Path,default=DEFAULT_ATTACHMENT)
    parser.add_argument("--template",type=Path,default=DEFAULT_TEMPLATE)
    parser.add_argument("--output-dir",type=Path,default=Q1_ROOT)
    return parser


def main(argv: list[str] | None = None) -> int:
    args=build_parser().parse_args(argv); output=args.output_dir.resolve(); output.mkdir(parents=True,exist_ok=True)
    data=load_attachment(args.attachment); result=solve_lp(data); checks=calculate_checks(result,data)
    if checks["mutual_exclusion_violation"] > TOL: result=solve_milp(data); checks=calculate_checks(result,data)
    mandatory=["balance_residual_max","state_residual_max","soc_bound_violation","charge_power_violation","discharge_power_violation","nonnegative_violation","curtailment_violation","mutual_exclusion_violation","terminal_soc_error"]
    failures={k:checks[k] for k in mandatory if checks[k] > TOL}
    if failures: raise RuntimeError(f"main checks failed: {failures}")
    independent=solve_independent(data); ind_checks=calculate_checks(independent,data)
    ind_failures={k:ind_checks[k] for k in mandatory if ind_checks[k] > TOL}
    if ind_failures: raise RuntimeError(f"independent checks failed: {ind_failures}")
    if abs(result["objective"]-independent["objective"]) > COST_TOL: raise RuntimeError("objective mismatch")
    export_trace(output/"q1_trace.csv",result,data); export_result(output/"result1.xlsx",args.template,result)
    metadata={"status":"optimal","method":result["method"],"objective":result["objective"],"independent_objective":independent["objective"],
              "objective_difference":abs(result["objective"]-independent["objective"]),"solver":"scipy","scipy_version":scipy.__version__,
              "attachment_sha256":sha256_file(args.attachment),"checks":checks,"independent_checks":ind_checks}
    (output/"q1_validation.json").write_text(json.dumps(metadata,ensure_ascii=False,indent=2,sort_keys=True),encoding="utf-8")
    print(f"method={result['method']}"); print(f"objective={result['objective']:.12f}"); print(f"independent_objective={independent['objective']:.12f}")
    print(output/"result1.xlsx"); return 0


if __name__ == "__main__":
    raise SystemExit(main())
