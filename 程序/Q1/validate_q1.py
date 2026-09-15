#!/usr/bin/env python3
"""Independent Q1 output validator."""

from __future__ import annotations

import argparse, csv, hashlib, json, math, re
from datetime import time as dt_time
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

T=144; DT=1/6; TOL=1e-7; COST_TOL=1e-6
PROJECT_ROOT=Path(__file__).resolve().parents[1]; Q1_ROOT=Path(__file__).resolve().parent
SPECIFIED={"10:00-10:10":61,"12:00-12:10":73,"14:00-14:10":85,"16:00-16:10":97,"18:00-18:10":109,"20:00-20:10":121}
COLUMNS=["t","source_timestamp","canonical_start","canonical_end","price","load_power","pv_power","load_energy","pv_energy","q","c","s","w","E_start","E_end","charge_power","discharge_power","purchase_cost"]

def sha256_file(path:Path)->str:
    h=hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()

def point_minutes(value:Any)->int:
    if isinstance(value,dt_time): return value.hour*60+value.minute
    text=str(value).strip().replace('：',':'); plus=text.endswith('+1')
    if plus: text=text[:-2]
    h,m=(int(x) for x in text.split(':',1)); total=h*60+m+(1440 if plus else 0)
    if total<0 or total>1440 or total%10: raise ValueError(f'invalid time: {value!r}')
    return total

def interval_minutes(value:Any)->tuple[int,int]:
    text=str(value).strip().replace('—','-').replace('–','-').replace('−','-').replace('：',':')
    parts=text.split('-')
    if len(parts)!=2: raise ValueError(f'invalid interval: {value!r}')
    start,end=point_minutes(parts[0]),point_minutes(parts[1])
    if end==0 and start>0: end=1440
    if end<=start: raise ValueError(f'invalid interval order: {value!r}')
    return start,end

def expected_interval(t:int)->tuple[int,int]: return (t-1)*10,t*10

def numeric(value:Any)->float:
    if isinstance(value,bool) or value is None: raise ValueError(f'not numeric: {value!r}')
    result=float(str(value).replace(',','')) if isinstance(value,str) else float(value)
    if not math.isfinite(result): raise ValueError('nonfinite')
    return result

def is_excel_number(value:Any)->bool: return isinstance(value,(int,float)) and not isinstance(value,bool)

def max_abs(values): return max((abs(v) for v in values),default=0.0)

def err(report,code,message,**ctx): report['errors'].append({'code':code,'message':message,**ctx})

def warn(report,code,message,**ctx): report['warnings'].append({'code':code,'message':message,**ctx})

def read_attachment(path:Path,report:dict[str,Any]):
    data={'price':[],'load_power':[],'pv_power':[]}
    try: wb=load_workbook(path,read_only=True,data_only=True)
    except Exception as exc: err(report,'attachment_open_failed',str(exc)); return data
    try:
        ws=wb['Sheet1']; rows=list(ws.iter_rows(min_row=2,max_row=145,values_only=True))
        for i,row in enumerate(rows,1):
            try:
                if point_minutes(row[0])!=i*10: err(report,'attachment_time',f'bad timestamp at t={i}')
                for key,value in zip(['price','load_power','pv_power'],row[1:4]): data[key].append(numeric(value))
            except Exception as exc: err(report,'attachment_value',str(exc),row=i+1)
    finally: wb.close()
    for key,values in data.items():
        if len(values)!=T: err(report,'attachment_length',f'{key} length {len(values)}')
        if key!='price' and any(v<0 for v in values): err(report,'attachment_negative',f'{key} negative')
    if len(data['price'])==T and any(v<=0 for v in data['price']): err(report,'attachment_price','prices must be positive')
    return data

def read_trace(path:Path,report:dict[str,Any]):
    result={k:[] for k in COLUMNS}
    with path.open('r',encoding='utf-8-sig',newline='') as stream:
        reader=csv.DictReader(stream)
        missing=set(COLUMNS)-set(reader.fieldnames or [])
        if missing: err(report,'trace_columns',f'missing {sorted(missing)}')
        for row in reader:
            for key in COLUMNS: result[key].append(row.get(key))
    if len(result['t'])!=T: err(report,'trace_length',f'rows {len(result["t"])}')
    return result

def validate_trace(trace,attachment,report,tolerance=TOL,cost_tolerance=COST_TOL):
    if len(trace['t'])!=T or any(len(attachment[k])!=T for k in attachment): return {'available':False}
    a={}
    try:
        for key in COLUMNS[4:]: a[key]=[numeric(v) for v in trace[key]]
    except Exception as exc: err(report,'trace_numeric',str(exc)); return {'available':False}
    comp=[]
    for i in range(T):
        comp += [abs(a['price'][i]-attachment['price'][i]),abs(a['load_power'][i]-attachment['load_power'][i]),abs(a['pv_power'][i]-attachment['pv_power'][i]),abs(a['load_energy'][i]-attachment['load_power'][i]*DT),abs(a['pv_energy'][i]-attachment['pv_power'][i]*DT)]
    if max_abs(comp)>tolerance: err(report,'trace_attachment','trace does not match attachment',residual=max_abs(comp))
    balance=[a['q'][i]+a['pv_energy'][i]+a['s'][i]-a['load_energy'][i]-a['c'][i]-a['w'][i] for i in range(T)]
    state=[a['E_end'][i]-(a['E_start'][i]+.9*a['c'][i]-a['s'][i]/.9) for i in range(T)]
    continuity=[a['E_start'][i]-a['E_end'][i-1] for i in range(1,T)]
    soc=a['E_start']+a['E_end']; power=max(0.0,max(a['charge_power'])-5000,max(a['discharge_power'])-5000)
    mutual=max(a['c'][i]*a['s'][i] for i in range(T)); terminal=max(abs(a['E_start'][0]-6000),abs(a['E_end'][-1]-6000))
    details={'available':True,'arrays':a,'balance_residual_max':max_abs(balance),'state_residual_max':max_abs(state),'continuity_residual_max':max_abs(continuity),
             'soc_bound_violation':max(0.0,1200-min(soc),max(soc)-10800),'power_bound_violation':power,'mutual_exclusion_violation':mutual,'terminal_soc_error':terminal,
             'objective_recomputed':sum(a['price'][i]*a['q'][i] for i in range(T)),'total_purchase':sum(a['q']),'total_charge':sum(a['c']),'total_discharge':sum(a['s']),'total_curtailment':sum(a['w']),
             'specified_purchase':{k:a['q'][v-1] for k,v in SPECIFIED.items()},
             'four_hour_blocks':{f'block_{b+1}':{'charge':sum(a['c'][b*24:(b+1)*24]),'discharge':sum(a['s'][b*24:(b+1)*24])} for b in range(6)}}
    for key,code in [('balance_residual_max','balance_residual'),('state_residual_max','state_residual'),('continuity_residual_max','state_continuity'),('soc_bound_violation','soc_bound'),('power_bound_violation','power_bound'),('mutual_exclusion_violation','mutual_exclusion'),('terminal_soc_error','terminal_soc')]:
        if details[key]>tolerance: err(report,code,f'{key} failed',value=details[key])
    return details


def validate_excel(path,trace_result,report,tolerance=TOL):
    details={'file_opened':False}; wb=load_workbook(path,data_only=True)
    try:
        details.update(file_opened=True,sheet_names=wb.sheetnames)
        if wb.sheetnames!=['计划购电量','充放电量']: err(report,'sheet_names','wrong sheet names',actual=wb.sheetnames)
        plan=wb['计划购电量']; storage=wb['充放电量']; details['dimensions']={'计划购电量':[plan.max_row,plan.max_column],'充放电量':[storage.max_row,storage.max_column]}
        if plan.max_row!=145 or plan.max_column!=2: err(report,'plan_dimensions','plan must be 145x2')
        q=[]
        for t in range(1,T+1):
            label,value=plan.cell(t+1,1).value,plan.cell(t+1,2).value
            try:
                if interval_minutes(label)!=expected_interval(t): err(report,'excel_time_mapping','wrong interval',row=t+1,label=label)
            except Exception as exc: err(report,'excel_time_mapping',str(exc),row=t+1)
            if not is_excel_number(value): err(report,'excel_non_numeric','purchase not numeric',row=t+1); q.append(math.nan)
            else: q.append(float(value))
        if trace_result.get('available'):
            residual=max_abs([q[i]-trace_result['arrays']['q'][i] for i in range(T)]); details['trace_comparison']={'q_residual_max':residual}
            if residual>tolerance: err(report,'excel_trace_q','purchase mismatch',residual=residual)
        if storage.max_row!=7 or storage.max_column!=5: err(report,'storage_dimensions','storage must be 7x5')
        if trace_result.get('available'):
            for b in range(6):
                ec=trace_result['four_hour_blocks'][f'block_{b+1}']['charge']; ed=trace_result['four_hour_blocks'][f'block_{b+1}']['discharge']
                ac,ad=storage.cell(b+2,2).value,storage.cell(b+2,3).value
                if not is_excel_number(ac) or abs(float(ac)-ec)>tolerance: err(report,'excel_charge','charge aggregate mismatch',block=b+1)
                if not is_excel_number(ad) or abs(float(ad)-ed)>tolerance: err(report,'excel_discharge','discharge aggregate mismatch',block=b+1)
        states={}
        for row in range(2,8):
            label,value=storage.cell(row,4).value,storage.cell(row,5).value
            if label is not None: states[point_minutes(label)]=value
        if states.get(0)!=6000 or states.get(1440)!=6000: err(report,'excel_terminal_soc','terminal SOC must be 6000',states=states)
        return details
    finally: wb.close()


def main(argv=None):
    p=argparse.ArgumentParser(); p.add_argument('--attachment',type=Path,default=PROJECT_ROOT/'附件'/'附件1.xlsx'); p.add_argument('--result',type=Path,default=Q1_ROOT/'result1.xlsx'); p.add_argument('--trace',type=Path,default=Q1_ROOT/'q1_trace.csv'); p.add_argument('--metadata',type=Path,default=Q1_ROOT/'q1_validation.json'); p.add_argument('--output',type=Path,default=Q1_ROOT/'q1_output_validation.json'); p.add_argument('--require-independent',action='store_true'); args=p.parse_args(argv)
    report={'pass':False,'errors':[],'warnings':[],'hashes':{}}
    for label,path in [('attachment',args.attachment),('result',args.result),('trace',args.trace),('metadata',args.metadata)]:
        if path.exists(): report['hashes'][label]=sha256_file(path)
        else: err(report,f'{label}_missing',f'{label} missing',path=str(path))
    attachment=read_attachment(args.attachment,report) if args.attachment.exists() else {'price':[],'load_power':[],'pv_power':[]}
    trace=read_trace(args.trace,report) if args.trace.exists() else {k:[] for k in COLUMNS}
    tr=validate_trace(trace,attachment,report); report['trace_validation']=tr
    if args.result.exists(): report['excel_validation']=validate_excel(args.result,tr,report)
    if args.metadata.exists():
        meta=json.loads(args.metadata.read_text(encoding='utf-8')); report['metadata']=meta
        if meta.get('status')!='optimal': err(report,'solver_status','status not optimal')
        for key in ['objective','independent_objective']:
            if meta.get(key) is None: err(report,f'{key}_missing',key)
            elif tr.get('available') and abs(float(meta[key])-tr['objective_recomputed'])>COST_TOL: err(report,f'{key}_mismatch',key)
    report['pass']=not report['errors']; args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2,sort_keys=True),encoding='utf-8')
    print(f"validation_pass={report['pass']} errors={len(report['errors'])} warnings={len(report['warnings'])}"); print(args.output.resolve()); return 0 if report['pass'] else 1

if __name__=='__main__': raise SystemExit(main())
