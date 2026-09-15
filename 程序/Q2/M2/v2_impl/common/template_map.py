#!/usr/bin/env python3
"""Q2 v2 frozen template mapping: plan sheet columns <-> 10-minute day slots.

Frozen in V2 spec: attachment time point t = interval [t-10min, t);
plan sheet col 2..144 = q[t-1] for t=2..144, col 145 = q[0] (circular day-first);
the '+1' suffix is a template label artifact.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

T = 144
COL_FIRST = 2
COL_LAST = 145
COL_TOTAL_KWH = 146
COL_TOTAL_COST = 147
PLAN_SHEET = '计划购电量'
N_TEMPLATE_ROWS = 335  # 1 header + 334 formal days

_LABEL_RE = re.compile(r'^\s*(\d{1,2}):(\d{1,2})\s*-\s*(\d{1,2}):(\d{1,2})\s*(\+1)?\s*$')


def slot_to_col(t: int) -> int:
    """Slot t (1-based; t=1 is the day's first interval [0:00,0:10)) -> plan column."""
    if not 1 <= t <= T:
        raise ValueError(f'slot {t} out of range [1,{T}]')
    return COL_LAST if t == 1 else t


def col_to_slot(col: int) -> int:
    if not COL_FIRST <= col <= COL_LAST:
        raise ValueError(f'column {col} out of range [{COL_FIRST},{COL_LAST}]')
    return 1 if col == COL_LAST else col


def parse_label(s: Any) -> tuple[int, int, int, int, bool] | None:
    if s is None:
        return None
    m = _LABEL_RE.match(str(s).strip())
    if not m:
        return None
    return (int(m.group(1)), int(m.group(2)), int(m.group(3)), int(m.group(4)),
            m.group(5) is not None)


def expected_label_semantics() -> list[tuple[int, int, int, int, bool]]:
    """Semantics of plan columns 2..145 in order: (h1, m1, h2, m2, plus1)."""
    out: list[tuple[int, int, int, int, bool]] = []
    for t in range(2, T + 1):
        lo, hi = 10 * (t - 1), 10 * t
        h2, m2 = divmod(hi % 1440, 60)
        out.append((lo // 60, lo % 60, h2, m2, hi >= 1440))
    out.append((0, 0, 0, 10, True))  # col 145 = day's first interval, circular
    return out


def label_text(sem: tuple[int, int, int, int, bool]) -> str:
    h1, m1, h2, m2, plus = sem
    return f'{h1}:{m1:02d}-{h2}:{m2:02d}' + ('+1' if plus else '')


def check_template(path: Path) -> list[str]:
    errs: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    if PLAN_SHEET not in wb.sheetnames:
        errs.append(f'template: sheet {PLAN_SHEET!r} missing')
        wb.close()
        return errs
    ws = wb[PLAN_SHEET]
    header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    want = expected_label_semantics()
    if len(header) < COL_LAST:
        errs.append(f'template: plan header width {len(header)} < {COL_LAST}')
    else:
        for i, sem in enumerate(want):
            got = header[COL_FIRST - 1 + i]
            if parse_label(got) != sem:
                errs.append(f'template: plan col {COL_FIRST + i} = {got!r} != {label_text(sem)!r}')
    for col, lab in ((COL_TOTAL_KWH, '全天购电量'), (COL_TOTAL_COST, '全天购电费')):
        if len(header) >= col and str(header[col - 1]).strip() != lab:
            errs.append(f'template: plan col {col} = {header[col - 1]!r} != {lab!r}')
    if ws.max_row != N_TEMPLATE_ROWS:
        errs.append(f'template: plan rows {ws.max_row} != {N_TEMPLATE_ROWS}')
    wb.close()
    return errs


def text_anomalies(path: Path) -> list[str]:
    """Cosmetic label deviations (semantics parse correctly, text differs).

    The template labels are machine-generated artifacts; these are informational.
    """
    out: list[str] = []
    wb = load_workbook(path, read_only=True, data_only=True)
    if PLAN_SHEET in wb.sheetnames:
        ws = wb[PLAN_SHEET]
        header = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        for i, sem in enumerate(expected_label_semantics()):
            got = header[COL_FIRST - 1 + i]
            if parse_label(got) == sem and str(got).strip() != label_text(sem):
                out.append(f'plan col {COL_FIRST + i}: {got!r} vs canonical {label_text(sem)!r}')
    wb.close()
    return out
