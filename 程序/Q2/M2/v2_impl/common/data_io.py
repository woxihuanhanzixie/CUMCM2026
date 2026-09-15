#!/usr/bin/env python3
"""Q2 v2 shared data I/O: read-only attachments and physical constants."""

from __future__ import annotations

import hashlib
from datetime import time as dt_time
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

T = 144
DT = 1 / 6
ETA_C = 0.9
ETA_D = 0.9
E_LO = 1200.0
E_HI = 10800.0
E_INIT = 6000.0
P_MAX = 5000.0
E_PORT_MAX = P_MAX * DT
RAMP_SLOPE = ETA_C * P_MAX * DT
EMERGENCY_MULT = 5.0
N_DAYS = 365
SCORED_START = 31
N_SCORED = N_DAYS - SCORED_START

PROJECT_ROOT = next(p for p in Path(__file__).resolve().parents
                    if (p / 'C题' / '附件' / '附件1.xlsx').is_file())
Q2_ROOT = Path(__file__).resolve().parents[2]
A1_DEFAULT = PROJECT_ROOT / 'C题' / '附件' / '附件1.xlsx'
A2_DEFAULT = PROJECT_ROOT / 'C题' / '附件' / '附件2.xlsx'
TEMPLATE_DEFAULT = PROJECT_ROOT / 'C题' / '附件' / '附件5' / 'result2.xlsx'

REF_HASH_A1 = '66b87134f5ecccd68184d3539bb1293ef039f9e0fdd955a589b9bfa7f227c377'
REF_HASH_A2 = '2e95fd446bfafa0d8c59577b5c2e2ea8b3f1def20dde54a3062556f4da9b4c72'


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def time_label(v: Any) -> str | None:
    if isinstance(v, dt_time):
        return f'{v.hour}:{v.minute:02d}'
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return None


def expected_a1_labels() -> list[str]:
    return [f'{m // 60}:{m % 60:02d}' for m in range(10, 1440, 10)] + ['0:00+1']


def expected_a2_header_labels() -> list[str]:
    return [f'{m // 60}:{m % 60:02d}' for m in range(10, 1440, 10)] + ['0:00+1']


def load_attachment1(path: Path = A1_DEFAULT) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, max_row=T + 1, values_only=True))
    wb.close()
    if len(rows) != T:
        raise ValueError(f'attachment1: expected {T} data rows, got {len(rows)}')
    labels = [time_label(r[0]) for r in rows]
    price = np.array([float(r[1]) for r in rows])
    load = np.array([float(r[2]) for r in rows])
    pv = np.array([float(r[3]) for r in rows])
    return {'labels': labels, 'price': price, 'load': load, 'pv': pv}


def load_attachment2(path: Path = A2_DEFAULT) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=True)
    load_ws = wb['小区负载']
    pv_ws = wb['光伏发电实际功率']
    load_header = next(load_ws.iter_rows(min_row=1, max_row=1, values_only=True))
    pv_header = next(pv_ws.iter_rows(min_row=1, max_row=1, values_only=True))
    load_rows = list(load_ws.iter_rows(min_row=2, values_only=True))
    pv_rows = list(pv_ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    dates = [r[0] for r in load_rows]
    pv_dates = [r[0] for r in pv_rows]
    load_act = np.array([[float(v) for v in r[1:T + 1]] for r in load_rows])
    pv_act = np.array([[float(v) for v in r[1:T + 1]] for r in pv_rows])
    if load_act.shape != (N_DAYS, T) or pv_act.shape != (N_DAYS, T):
        raise ValueError(f'attachment2: shape {(load_act.shape, pv_act.shape)} != {(N_DAYS, T)}')
    if dates != pv_dates:
        raise ValueError('attachment2: date columns differ between sheets')
    return {'dates': dates, 'load_act': load_act, 'pv_act': pv_act,
            'load_header': load_header, 'pv_header': pv_header}


def check_attachment2_header(load_header, pv_header) -> list[str]:
    errs: list[str] = []
    want = expected_a2_header_labels()
    for name, header in (('load', load_header), ('pv', pv_header)):
        if len(header) != T + 1:
            errs.append(f'attachment2 {name}: header width {len(header)} != {T + 1}')
            continue
        for i, (cell, w) in enumerate(zip(header[1:], want)):
            got = time_label(cell)
            if got != w:
                errs.append(f'attachment2 {name}: header col {i + 2} = {cell!r} != {w!r}')
    return errs
