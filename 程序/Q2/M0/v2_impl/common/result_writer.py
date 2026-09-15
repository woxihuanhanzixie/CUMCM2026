#!/usr/bin/env python3
"""Q2 v2 shared result2 workbook writer + readback (frozen template mapping, V2 §1).

All three models write through this module so that the plan-sheet column
mapping, the six 4-hour charge/discharge groups and the emergency-interval
rows are identical across candidates.

Plan sheet: col 2..144 = slots 2..144, col 145 = slot 1 (circular day-first);
col 146 = 全天购电量 (sum q), col 147 = 全天购电费 (sum p*q).
Charge sheet: 6 rows per day, groups 0:00-4:00 ... 20:00-24:00, 储电量 = E at
group end; 时刻 column replicates the template quirk (00:00:00 on the first
group row, 24:00 on the second).
Emergency sheet: one row per contiguous run of slots with z > 0, interval
label 'H:MM-H:MM', 购电量 = sum z; date only on the first row of each day.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook

from .data_io import T
from .template_map import (COL_FIRST, COL_LAST, COL_TOTAL_KWH, COL_TOTAL_COST,
                           PLAN_SHEET, expected_label_semantics, label_text,
                           slot_to_col)

CHARGE_SHEET = '充放电量'
EMERGENCY_SHEET = '紧急购电量'
GROUP_LABELS = ['0:00-4:00', '4:00-8:00', '8:00-12:00',
                '12:00-16:00', '16:00-20:00', '20:00-24:00']
EMG_TOL = 1e-9
VALUE_TOL = 1e-9
KWH_TOL = 1e-6
COST_TOL = 1e-3


def plan_header() -> list[str]:
    hdr = ['日期\\时间']
    for sem in expected_label_semantics():
        hdr.append(label_text(sem))
    hdr += ['全天购电量', '全天购电费']
    return hdr


def interval_label(t0: int, t1: int) -> str:
    """Contiguous 0-based slot run [t0, t1] -> 'H:MM-H:MM' label."""
    lo, hi = 10 * t0, 10 * (t1 + 1)
    return f'{lo // 60}:{lo % 60:02d}-{hi // 60}:{hi % 60:02d}'


def parse_interval_label(s: str) -> tuple[int, int] | None:
    """'H:MM-H:MM' -> (lo_min, hi_min) in [0, 1440], or None if malformed."""
    try:
        a, b = str(s).strip().split('-')
        h1, m1 = a.split(':')
        h2, m2 = b.split(':')
        lo, hi = int(h1) * 60 + int(m1), int(h2) * 60 + int(m2)
    except (ValueError, AttributeError):
        return None
    if not (0 <= lo < hi <= 1440) or lo % 10 or hi % 10:
        return None
    return lo, hi


def write_result2(path: Path, dates, q, c, s, z, E, price) -> None:
    q = np.asarray(q, float)
    c = np.asarray(c, float)
    s = np.asarray(s, float)
    z = np.asarray(z, float)
    E = np.asarray(E, float)
    price = np.asarray(price, float)
    n_days = q.shape[0]
    if not all(a.shape == (n_days, T) for a in (q, c, s, z, E)):
        raise ValueError('write_result2: q/c/s/z/E must all be (n_days, 144)')

    wb = Workbook()
    ws1 = wb.active
    ws1.title = PLAN_SHEET
    ws1.append(plan_header())
    for d in range(n_days):
        row = [dates[d]]
        for col in range(COL_FIRST, COL_LAST + 1):
            t = 1 if col == COL_LAST else col   # col 145 = slot 1 (circular)
            row.append(float(q[d, t - 1]))
        row.append(float(q[d].sum()))
        row.append(float(np.dot(price, q[d])))
        ws1.append(row)

    ws2 = wb.create_sheet(CHARGE_SHEET)
    ws2.append(['日期', '时间段', '充电量', '放电量', '时刻', '储电量'])
    for d in range(n_days):
        for g in range(6):
            sl = slice(g * 24, (g + 1) * 24)
            ws2.append([dates[d] if g == 0 else None,
                        GROUP_LABELS[g],
                        float(c[d, sl].sum()),
                        float(s[d, sl].sum()),
                        '00:00:00' if g == 0 else ('24:00' if g == 1 else None),
                        float(E[d, (g + 1) * 24 - 1])])

    ws3 = wb.create_sheet(EMERGENCY_SHEET)
    ws3.append(['日期', '购电时间段', '购电量'])
    for d in range(n_days):
        first = True
        t0 = 0
        while t0 < T:
            if z[d, t0] <= EMG_TOL:
                t0 += 1
                continue
            t1 = t0
            while t1 + 1 < T and z[d, t1 + 1] > EMG_TOL:
                t1 += 1
            ws3.append([dates[d] if first else None,
                        interval_label(t0, t1),
                        float(z[d, t0:t1 + 1].sum())])
            first = False
            t0 = t1 + 1
    wb.save(path)


def check_result2(path: Path, dates, q, c, s, z, E, price) -> list[str]:
    """Read back a candidate workbook and verify it against the exact arrays."""
    errs: list[str] = []
    q = np.asarray(q, float)
    c = np.asarray(c, float)
    s = np.asarray(s, float)
    z = np.asarray(z, float)
    E = np.asarray(E, float)
    price = np.asarray(price, float)
    n_days = q.shape[0]
    wb = load_workbook(path, read_only=True, data_only=True)

    if wb.sheetnames != [PLAN_SHEET, CHARGE_SHEET, EMERGENCY_SHEET]:
        errs.append(f'workbook sheets {wb.sheetnames} != expected order')

    ws1 = wb[PLAN_SHEET]
    rows1 = list(ws1.iter_rows(values_only=True))
    if len(rows1) != n_days + 1:
        errs.append(f'plan sheet rows {len(rows1)} != {n_days + 1}')
    else:
        header = [str(v) for v in rows1[0]]
        if header != plan_header():
            errs.append('plan sheet header mismatch')
        for d in range(n_days):
            row = rows1[d + 1]
            if row[0] != dates[d]:
                errs.append(f'plan sheet day {d} date {row[0]!r} != {dates[d]!r}')
                break
            for t in range(1, T + 1):
                if abs(float(row[slot_to_col(t) - 1]) - q[d, t - 1]) > VALUE_TOL:
                    errs.append(f'plan sheet day {d} slot {t}: '
                                f'{row[slot_to_col(t) - 1]} != {q[d, t - 1]}')
                    break
            if abs(float(row[COL_TOTAL_KWH - 1]) - q[d].sum()) > KWH_TOL:
                errs.append(f'plan sheet day {d} total kWh mismatch')
            if abs(float(row[COL_TOTAL_COST - 1]) - np.dot(price, q[d])) > COST_TOL:
                errs.append(f'plan sheet day {d} total cost mismatch')

    ws2 = wb[CHARGE_SHEET]
    rows2 = list(ws2.iter_rows(values_only=True))
    if len(rows2) != n_days * 6 + 1:
        errs.append(f'charge sheet rows {len(rows2)} != {n_days * 6 + 1}')
    else:
        for d in range(n_days):
            for g in range(6):
                row = rows2[1 + d * 6 + g]
                sl = slice(g * 24, (g + 1) * 24)
                if row[1] != GROUP_LABELS[g]:
                    errs.append(f'charge sheet day {d} group {g} label {row[1]!r}')
                if abs(float(row[2]) - c[d, sl].sum()) > KWH_TOL:
                    errs.append(f'charge sheet day {d} group {g} charge sum mismatch')
                if abs(float(row[3]) - s[d, sl].sum()) > KWH_TOL:
                    errs.append(f'charge sheet day {d} group {g} discharge sum mismatch')
                if abs(float(row[5]) - E[d, (g + 1) * 24 - 1]) > KWH_TOL:
                    errs.append(f'charge sheet day {d} group {g} E mismatch')
                if (row[0] == dates[d]) != (g == 0):
                    errs.append(f'charge sheet day {d} group {g} date cell rule violated')

    ws3 = wb[EMERGENCY_SHEET]
    intervals: dict[int, list[tuple[tuple[int, int], float]]] = {}
    cur_day: int | None = None
    for row in ws3.iter_rows(min_row=2, values_only=True):
        date_cell, label_cell, amt_cell = row[0], row[1], row[2]
        if date_cell is not None:
            try:
                cur_day = dates.index(date_cell)
            except ValueError:
                errs.append(f'emergency sheet unknown date {date_cell!r}')
                cur_day = None
                continue
        if cur_day is None:
            continue
        parsed = parse_interval_label(str(label_cell))
        if parsed is None:
            errs.append(f'emergency sheet day {cur_day} bad interval {label_cell!r}')
            continue
        intervals.setdefault(cur_day, []).append((parsed, float(amt_cell)))
    for d in range(n_days):
        listed = intervals.get(d, [])
        mask = np.zeros(T, bool)
        for (lo, hi), _amt in listed:
            mask[lo // 10:hi // 10] = True
        want = z[d] > EMG_TOL
        if not np.array_equal(mask, want):
            errs.append(f'emergency sheet day {d} interval coverage mismatch')
            continue
        for i, ((lo, hi), amt) in enumerate(listed):
            if abs(amt - z[d, lo // 10:hi // 10].sum()) > KWH_TOL:
                errs.append(f'emergency sheet day {d} interval {i} amount mismatch')
    wb.close()
    return errs
